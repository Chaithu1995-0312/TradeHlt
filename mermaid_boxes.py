import openpyxl
from openpyxl.styles import Font, PatternFill

# Create a new workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Mermaid Boxes"

# Add header
ws['A1'] = "Box Name"
ws['A1'].font = Font(bold=True)
ws['A1'].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
ws['A1'].font = Font(bold=True, color="FFFFFF")

# Add box names from L0 diagram
boxes = [
    "agent",
    "analytics",
    "bitnet",
    "charts",
    "cognitive",
    "config_layer",
    "control_plane",
    "core",
    "data_ingestion",
    "engines",
    "events",
    "execution",
    "expansion",
    "features",
    "feedback",
    "governance",
    "identity",
    "inout",
    "interpreters",
    "journal",
    "live",
    "llm_research",
    "monitoring",
    "msip",
    "multi_llm",
    "portfolio",
    "regime",
    "replay",
    "research",
    "retrieval",
    "runtime",
    "scanner",
    "search",
    "strategies",
    "structure",
    "training",
    "uat",
    "ui",
    "utils",
    "validation_access"
]

# Write boxes to Excel
for idx, box in enumerate(boxes, start=2):
    ws[f'A{idx}'] = box

# Adjust column width
ws.column_dimensions['A'].width = 25

# Save the file
wb.save('D:\\Tradelatest\\mermaid_boxes.xlsx')
print(f"Excel file created with {len(boxes)} boxes")
