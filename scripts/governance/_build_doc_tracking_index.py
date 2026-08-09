"""Build DOC_TRACKING_INDEX.xlsx — metadata-only inventory (no content reads)."""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"

TOPIC_MAP = {
    "architecture": "Architecture & Flow",
    "governance": "Governance & Authority",
    "reference": "Reference (Schemas/CLI/Config)",
    "topics": "Topic Index (Concept↔Code)",
    "analysis": "Analysis (Point-in-time)",
    "implementation_plan": "Implementation Plans",
    "plans": "Plans",
    "research": "Research Programs",
    "research-readiness": "Research Readiness",
    "operations": "Operations",
    "control_plane": "Control Plane",
    "handover": "Handover / Transfer",
    "intent": "Intent / MIAR",
    "human-language-analysis": "Human-Language Analysis",
    "findings": "Findings & Knowledge",
}

ROOT_TOPIC_RULES = [
    (r"^ZONE-X", "Zone-X Program"),
    (r"^CLAUDE|^AGENTS", "Operating Manual"),
    (r"^assistant_project|^llm_project", "Session Logs"),
    (r"^HANDOFF|^README", "Entry / Handoff"),
    (
        r"^YAML_CONSUMER|^CODEBASE_WIRING|^COMPLETE_CODEBASE|^PHASE_6|^MASTER_ARCH|^CODE_ANALYSIS",
        "Wiring / Atlas Audits",
    ),
    (r"^AMBIGUITY|^RESOLUTION|^FALLBACK|^CLEANUP", "Ambiguity / Resolution Registry"),
    (r"^preregistered", "Research Preregistration"),
]

SESSION_DATE = datetime(2026, 8, 6)


def infer_topic(rel: Path) -> str:
    parts = rel.parts
    if len(parts) >= 2 and parts[0] == "docs":
        seg = parts[1].lower()
        if seg in TOPIC_MAP:
            return TOPIC_MAP[seg]
        return f"Docs / {parts[1]}"
    if len(parts) >= 1 and parts[0] == "multi_llm":
        return "Multi-LLM Workflow"
    if len(parts) >= 1 and parts[0].lower().startswith("chatgpt"):
        return "Multi-LLM Workflow"
    name = rel.name
    for pat, topic in ROOT_TOPIC_RULES:
        if re.search(pat, name, re.I):
            return topic
    if rel.suffix.lower() == ".md" and len(parts) == 1:
        return "Root Misc"
    return "Other"


def authority_hint(rel: Path, name: str) -> str:
    n = name.lower()
    p = str(rel).replace("\\", "/").lower()
    if name in ("CLAUDE.md", "AGENTS.md"):
        return "TIER0_OPERATING"
    if "current-findings" in n:
        return "TIER0_CONCLUSIONS"
    if "active_models" in n:
        return "TIER0_MODELS"
    if p.startswith("docs/reference/") or p.startswith("docs/architecture/"):
        return "TIER1_LIVING"
    if p.startswith("docs/governance/"):
        return "TIER1_GOVERNANCE"
    if p.startswith("docs/topics/"):
        return "TIER1_TOPICS"
    if "closure" in n or "contract" in n:
        return "TIER1_CLOSURE"
    if p.startswith("docs/analysis/") or "archived" in p:
        return "TIER3_POINT_IN_TIME"
    if p.startswith("docs/implementation_plan/") or p.startswith("docs/plans/"):
        return "TIER2_PLANS"
    if name.startswith("ZONE-X"):
        return "TIER2_PROGRAM"
    if "assistant_project" in n or "handoff" in n.lower():
        return "TIER2_SESSION"
    if "generated" in n:
        return "GENERATED"
    if p.startswith("multi_llm/") or "chatgpt" in p:
        return "TIER2_WORKFLOW"
    return "UNCLASSIFIED"


def freshness_bucket(dt: datetime) -> str:
    days = (SESSION_DATE - dt.replace(tzinfo=None)).days
    if days <= 1:
        return "TODAY/YESTERDAY"
    if days <= 7:
        return "LAST_7D"
    if days <= 30:
        return "LAST_30D"
    if days <= 90:
        return "LAST_90D"
    return "OLDER_90D+"


def family_key(name: str) -> str:
    s = re.sub(r"[-_]v?\d+(\.\d+)*", "", name, flags=re.I)
    s = re.sub(r"[-_]20\d{2}[-_]\d{2}[-_]\d{2}", "", s)
    s = re.sub(r"[-_]\d{8}", "", s)
    s = re.sub(r"\.md$", "", s, flags=re.I)
    s = re.sub(r"[-_]+", "-", s).strip("-").lower()
    return s or name.lower()


def collect_files() -> list[Path]:
    files: list[Path] = []
    if DOCS.exists():
        files.extend(p for p in DOCS.rglob("*.md") if p.is_file())
    files.extend(ROOT.glob("*.md"))
    ml = ROOT / "multi_llm"
    if ml.exists():
        files.extend(p for p in ml.rglob("*.md") if p.is_file())
    for cg_name in ("ChatGpt workflow", "ChatGpt Workflow"):
        cg = ROOT / cg_name
        if cg.exists():
            files.extend(p for p in cg.rglob("*.md") if p.is_file())
            break
    return files


def main() -> None:
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    primary_fill = PatternFill("solid", fgColor="C6EFCE")
    secondary_fill = PatternFill("solid", fgColor="FFEB9C")
    archive_fill = PatternFill("solid", fgColor="F4B183")
    today_fill = PatternFill("solid", fgColor="BDD7EE")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    seen: set[Path] = set()
    rows: list[dict] = []
    for p in collect_files():
        try:
            rp = p.resolve()
        except OSError:
            rp = p
        if rp in seen:
            continue
        seen.add(rp)
        try:
            st = p.stat()
            mtime = datetime.fromtimestamp(st.st_mtime)
            size = st.st_size
        except OSError:
            continue
        try:
            rel = p.relative_to(ROOT)
        except ValueError:
            rel = Path(p.name)
        topic = infer_topic(rel)
        sub = rel.parent.name if rel.parent != Path(".") else "root"
        rows.append(
            {
                "topic": topic,
                "subfolder": sub,
                "doc_name": p.name,
                "rel_path": str(rel).replace("\\", "/"),
                "mtime": mtime,
                "mtime_str": mtime.strftime("%Y-%m-%d %H:%M"),
                "date": mtime.strftime("%Y-%m-%d"),
                "size_kb": round(size / 1024, 1),
                "authority": authority_hint(rel, p.name),
                "freshness": freshness_bucket(mtime),
                "status": "NOT_READ",
                "notes": "",
                "latest_in_family": "",
                "ref_priority": "",
            }
        )

    by_family: dict[tuple[str, str], list[int]] = defaultdict(list)
    for i, r in enumerate(rows):
        by_family[(r["topic"], family_key(r["doc_name"]))].append(i)

    for _key, idxs in by_family.items():
        if len(idxs) <= 1:
            rows[idxs[0]]["latest_in_family"] = "ONLY"
            rows[idxs[0]]["ref_priority"] = "PRIMARY"
            continue
        idxs_sorted = sorted(idxs, key=lambda i: rows[i]["mtime"], reverse=True)
        latest = idxs_sorted[0]
        for rank, i in enumerate(idxs_sorted):
            if i == latest:
                rows[i]["latest_in_family"] = "LATEST"
                rows[i]["ref_priority"] = "PRIMARY"
                rows[i]["notes"] = f"Family has {len(idxs)} docs; use this on ambiguity"
            else:
                rows[i]["latest_in_family"] = f"OLDER_rank{rank + 1}"
                rows[i]["ref_priority"] = "SECONDARY" if rank == 1 else "ARCHIVE_CANDIDATE"
                rows[i]["notes"] = (
                    f"Older than {rows[latest]['doc_name']} ({rows[latest]['mtime_str']})"
                )

    rows.sort(key=lambda r: (r["topic"], -r["mtime"].timestamp(), r["doc_name"]))

    wb = Workbook()

    # Sheet 1: Master Index
    ws = wb.active
    ws.title = "Master_Index"
    headers = [
        "Topic",
        "Subfolder",
        "Doc_Name",
        "Rel_Path",
        "Last_Modified",
        "Date",
        "Size_KB",
        "Authority_Tier",
        "Freshness",
        "Latest_In_Family",
        "Ref_Priority",
        "Read_Status",
        "Discussion_Notes",
        "User_Intent_Tag",
        "Ambiguity_Flag",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for r in rows:
        if r["latest_in_family"] not in ("ONLY", "LATEST"):
            amb = "YES"
        elif r["latest_in_family"] == "LATEST":
            amb = "FAMILY_LATEST"
        else:
            amb = ""
        ws.append(
            [
                r["topic"],
                r["subfolder"],
                r["doc_name"],
                r["rel_path"],
                r["mtime_str"],
                r["date"],
                r["size_kb"],
                r["authority"],
                r["freshness"],
                r["latest_in_family"],
                r["ref_priority"],
                r["status"],
                r["notes"],
                "",
                amb,
            ]
        )
        row_idx = ws.max_row
        pri = r["ref_priority"]
        fill = None
        if pri == "PRIMARY" and r["freshness"] in ("TODAY/YESTERDAY", "LAST_7D"):
            fill = today_fill
        elif pri == "PRIMARY":
            fill = primary_fill
        elif pri == "SECONDARY":
            fill = secondary_fill
        elif pri == "ARCHIVE_CANDIDATE":
            fill = archive_fill
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row_idx, c)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if fill and c in (10, 11):
                cell.fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    widths = [28, 18, 48, 72, 18, 12, 10, 18, 14, 14, 16, 12, 40, 18, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 2: Topic Summary
    ws2 = wb.create_sheet("Topic_Summary")
    ws2.append(
        [
            "Topic",
            "Doc_Count",
            "Latest_Doc",
            "Latest_Date",
            "Oldest_Date",
            "Primary_Count",
            "Archive_Candidates",
            "Suggested_Entry_Docs",
        ]
    )
    for c in range(1, 9):
        cell = ws2.cell(1, c)
        cell.font = header_font
        cell.fill = header_fill

    by_topic: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        by_topic[r["topic"]].append(r)

    for topic in sorted(by_topic.keys()):
        items = by_topic[topic]
        items_s = sorted(items, key=lambda x: x["mtime"], reverse=True)
        primaries = [x for x in items if x["ref_priority"] == "PRIMARY"]
        archives = [x for x in items if x["ref_priority"] == "ARCHIVE_CANDIDATE"]
        entry = sorted(primaries, key=lambda x: x["mtime"], reverse=True)[:3]
        entry_str = " | ".join(e["doc_name"] for e in entry)
        ws2.append(
            [
                topic,
                len(items),
                items_s[0]["doc_name"],
                items_s[0]["date"],
                min(x["date"] for x in items),
                len(primaries),
                len(archives),
                entry_str,
            ]
        )
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:H{ws2.max_row}"
    for i, w in enumerate([32, 12, 48, 12, 12, 14, 16, 80], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Sheet 3: Ambiguity families
    ws3 = wb.create_sheet("Ambiguity_Families")
    ws3.append(
        [
            "Topic",
            "Family_Key",
            "Count",
            "LATEST_Doc",
            "LATEST_Path",
            "LATEST_Date",
            "Older_Docs",
            "Rule",
        ]
    )
    red_fill = PatternFill("solid", fgColor="C00000")
    for c in range(1, 9):
        cell = ws3.cell(1, c)
        cell.font = header_font
        cell.fill = red_fill

    family_groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for r in rows:
        family_groups[(r["topic"], family_key(r["doc_name"]))].append(r)

    amb_count = 0
    for (topic, fam), items in sorted(family_groups.items()):
        if len(items) < 2:
            continue
        amb_count += 1
        items_s = sorted(items, key=lambda x: x["mtime"], reverse=True)
        latest = items_s[0]
        older = "; ".join(f"{x['doc_name']}@{x['date']}" for x in items_s[1:])
        ws3.append(
            [
                topic,
                fam,
                len(items),
                latest["doc_name"],
                latest["rel_path"],
                latest["date"],
                older,
                "Prefer LATEST mtime on ambiguity; verify Authority_Tier before treating as truth",
            ]
        )
    ws3.freeze_panes = "A2"
    if ws3.max_row >= 1:
        ws3.auto_filter.ref = f"A1:H{ws3.max_row}"
    for i, w in enumerate([28, 36, 8, 40, 60, 12, 60, 50], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # Sheet 4: Hot recent
    ws4 = wb.create_sheet("Hot_Recent_30d")
    ws4.append(
        [
            "Topic",
            "Doc_Name",
            "Rel_Path",
            "Last_Modified",
            "Authority_Tier",
            "Ref_Priority",
            "Why_Hot",
        ]
    )
    blue_fill = PatternFill("solid", fgColor="2E75B6")
    for c in range(1, 8):
        cell = ws4.cell(1, c)
        cell.font = header_font
        cell.fill = blue_fill
    hot = [r for r in rows if r["freshness"] in ("TODAY/YESTERDAY", "LAST_7D", "LAST_30D")]
    hot.sort(key=lambda x: x["mtime"], reverse=True)
    for r in hot:
        why = "mtime-bucket " + r["freshness"]
        if r["authority"].startswith("TIER0"):
            why += "; operating authority"
        ws4.append(
            [
                r["topic"],
                r["doc_name"],
                r["rel_path"],
                r["mtime_str"],
                r["authority"],
                r["ref_priority"],
                why,
            ]
        )
    ws4.freeze_panes = "A2"
    if ws4.max_row >= 1:
        ws4.auto_filter.ref = f"A1:G{ws4.max_row}"
    for i, w in enumerate([28, 48, 72, 18, 18, 14, 40], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    # Sheet 5: Discussion tracker
    ws5 = wb.create_sheet("Discussion_Tracker")
    ws5.append(
        [
            "Turn",
            "Date",
            "User_Intent_Slice",
            "Topics_Engaged",
            "Docs_Referenced",
            "Docs_Read_This_Turn",
            "Decision",
            "Open_Question",
            "Next_Probe",
            "Status",
        ]
    )
    green_fill = PatternFill("solid", fgColor="548235")
    for c in range(1, 11):
        cell = ws5.cell(1, c)
        cell.font = header_font
        cell.fill = green_fill
    ws5.append(
        [
            1,
            "2026-08-06",
            "Pre-intent: inventory + role selection",
            "ALL (index only)",
            "NONE (not read)",
            "NONE",
            "Tracking workbook created; role=Architect/Facilitator",
            "Await full user intent",
            "Receive intent → map to topics/docs → propose plan discussion agenda",
            "WAITING_INTENT",
        ]
    )
    ws5.freeze_panes = "A2"
    for i, w in enumerate([8, 12, 36, 28, 40, 28, 40, 36, 40, 16], 1):
        ws5.column_dimensions[get_column_letter(i)].width = w

    # Sheet 6: Role protocol
    ws6 = wb.create_sheet("Role_Protocol")
    ws6.append(["Field", "Value"])
    for c in range(1, 3):
        cell = ws6.cell(1, c)
        cell.font = header_font
        cell.fill = header_fill
    protocol = [
        (
            "Selected_Role",
            "Architect / Discussion Facilitator (ChatGPT-Interpreter analogue in multi_llm)",
        ),
        (
            "Why_This_Role",
            "Plan discussion + full intent incoming; structure topics, map docs, surface conflicts — do NOT execute code or deep-read yet",
        ),
        (
            "Not_Doing",
            "DeepSeek Planner decomposition until intent lands; Claude Executor code; research claims",
        ),
        (
            "Ambiguity_Rule",
            "On conflicting docs: prefer highest Authority_Tier; within same tier prefer Latest_In_Family=LATEST by mtime; if still ambiguous → surface TruthConflict, do not invent",
        ),
        (
            "Read_Policy",
            "NOT_READ until user intent names a slice; then read only PRIMARY docs for engaged topics",
        ),
        (
            "Authority_Precedence",
            "TIER0 (CLAUDE/findings/ACTIVE_VERSION) > TIER1 living > TIER2 plans/session > TIER3 analysis point-in-time > GENERATED",
        ),
        ("User_Bridge", "User remains principal decider; this role recommends and tracks"),
        (
            "Workbook_Path",
            "docs/governance/DOC_TRACKING_INDEX.xlsx (+ root copy DOC_TRACKING_INDEX.xlsx)",
        ),
        (
            "Update_Rule",
            "Each discussion turn: append Discussion_Tracker row; flip Read_Status when docs are opened",
        ),
        (
            "Proactive_Stance",
            "Before each intent slice: propose 3–5 entry docs from Topic_Summary + Hot_Recent; ask 1–2 sharp probes",
        ),
    ]
    for k, v in protocol:
        ws6.append([k, v])
    ws6.column_dimensions["A"].width = 24
    ws6.column_dimensions["B"].width = 110
    for row in ws6.iter_rows(min_row=2, max_row=ws6.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin

    # Sheet 7: Legend
    ws7 = wb.create_sheet("Legend")
    ws7.append(["Token", "Meaning"])
    ws7["A1"].font = header_font
    ws7["A1"].fill = header_fill
    ws7["B1"].font = header_font
    ws7["B1"].fill = header_fill
    legend = [
        ("PRIMARY", "Preferred reference for its family/topic (latest or only)"),
        ("SECONDARY", "Valid but older sibling — use only if PRIMARY insufficient"),
        ("ARCHIVE_CANDIDATE", "Superseded by newer family member — history only"),
        ("TIER0_*", "Operating truth / conclusions / model registry — do not silently override"),
        ("TIER1_*", "Living docs (architecture, governance, topics, closures)"),
        ("TIER2_*", "Plans, programs, session logs — useful, not runtime truth"),
        (
            "TIER3_POINT_IN_TIME",
            "Historical analysis — never treat as current state without revalidation",
        ),
        ("GENERATED", "Derived artifact — edit source, not this file"),
        ("NOT_READ", "Content not yet loaded this conversation"),
        ("Freshness buckets", "Relative to 2026-08-06 session date"),
        ("Green priority cells", "PRIMARY"),
        ("Yellow", "SECONDARY"),
        ("Orange", "ARCHIVE_CANDIDATE"),
        ("Blue priority cells", "PRIMARY + recent (7d)"),
    ]
    for a, b in legend:
        ws7.append([a, b])
    ws7.column_dimensions["A"].width = 24
    ws7.column_dimensions["B"].width = 90

    out = ROOT / "docs" / "governance" / "DOC_TRACKING_INDEX.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    out2 = ROOT / "DOC_TRACKING_INDEX.xlsx"
    wb.save(out2)

    print(f"TOTAL_ROWS={len(rows)}")
    print(f"TOPICS={len(by_topic)}")
    print(f"AMBIGUITY_FAMILIES={amb_count}")
    print(f"HOT_30D={len(hot)}")
    print(f"SAVED={out}")
    print(f"SAVED2={out2}")
    print("---TOPIC_COUNTS---")
    for t in sorted(by_topic.keys()):
        print(f"  {t}: {len(by_topic[t])}")
    print("---TOP20_HOT---")
    for r in hot[:20]:
        print(f"  {r['date']} | {r['topic'][:28]:28} | {r['doc_name']}")


if __name__ == "__main__":
    main()
