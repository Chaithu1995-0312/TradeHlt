#!/usr/bin/env python3
"""
Build an Excel inventory of tests/**/*.py business functionality.

Columns:
  - File Name
  - Summary of functionality
  - Referred files
  - Test Class
  - Intent of test class  (combined method intents)

Read-only analysis. No nested dependency deep-walk — only direct imports
resolved to project-style paths when possible.
"""
from __future__ import annotations

import argparse
import ast
import re
import subprocess
import sys
from pathlib import Path
from typing import Iterable, NamedTuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as e:  # pragma: no cover
    raise SystemExit(f"openpyxl required: {e}") from e

REPO = Path(__file__).resolve().parents[2]
TESTS = REPO / "tests"
OUT = REPO / "docs" / "analysis" / "tests_functionality_inventory.xlsx"
GROK_DIR = TESTS / "Grok"
GROK_OUT = REPO / "docs" / "analysis" / "grok_test_intent.xlsx"
CLAUDE_DIR = TESTS / "Claude"
CLAUDE_OUT = REPO / "docs" / "analysis" / "claude_test_intent.xlsx"


class Suite(NamedTuple):
    """One auditor suite: where its tests live and where its intent book goes."""

    name: str           # "Grok" | "Claude" — also the `Added by` column value
    directory: Path     # tests/<name>
    out: Path           # docs/analysis/<name.lower()>_test_intent.xlsx
    sheet: str          # primary sheet title
    pointer_sheet: str  # pointer sheet title on the class-grain book


_SUITES: dict[str, Suite] = {
    "Grok": Suite("Grok", GROK_DIR, GROK_OUT, "Grok Method Intent", "Grok Intent Pointer"),
    "Claude": Suite(
        "Claude", CLAUDE_DIR, CLAUDE_OUT, "Claude Method Intent", "Claude Intent Pointer"
    ),
}


def _suite(suite: str | Suite) -> Suite:
    return suite if isinstance(suite, Suite) else _SUITES[suite]


# Family letters are globally unique across suites: Grok owns A-I, Claude owns
# J-M. A letter absent from this map is still inventoried, with Topic layer
# UNMAPPED, so a new family never needs a design-doc edit to appear.
_FAMILY_LAYER = {
    # ── Grok (tests/Grok, families A-I) ──
    "A": "2 CRT 9-state M15 spine",
    "B": "8 Execution geometry + Ultron",
    "C": "7 Scoring / fusion / decision",
    "D": "9 Governance / promotion / measurement",
    "E": "7 Scoring / fusion / decision",
    "F": "1 Ontology / feature vector",
    "G": "0 OHLCV + clocks + no-lookahead",
    "H": "0 OHLCV + clocks + no-lookahead",
    "I": "4 Parent 3-candle 12-state",
    # ── Claude (tests/Claude, families J-M) ──
    "J": "3 Directional displacement F-074",
    "K": "6 SMC primitives F-076",
    "L": "1 Ontology / feature vector",
    "M": "4 Parent 3-candle 12-state",
}

_GROK_FAMILY_LAYER = _FAMILY_LAYER   # back-compat alias

_GROK_LABELS = (
    "Source:",
    "Failure mode:",
    "Why ordinary tests miss it:",
    "Numbers:",
)

_INTENT_HEADERS = [
    "Family",
    "Topic layer",
    "File",
    "Method",
    "Nodeid",
    "Case id",
    "Intent",
    "Trying to do",
    "Source contract",
    "Failure mode",
    "Why ordinary tests miss it",
    "Intent quality",
    "Parametrize arity",
    "Added by",
    "Referred files",
]


_GROK_INTENT_HEADERS = _INTENT_HEADERS   # back-compat alias


class CollectFailed(RuntimeError):
    """pytest --collect-only on an auditor suite failed or timed out."""

# stdlib / third-party prefixes to keep as package names (not project files)
_NON_PROJECT_ROOTS = frozenset(
    {
        "abc",
        "argparse",
        "ast",
        "asyncio",
        "base64",
        "builtins",
        "collections",
        "concurrent",
        "contextlib",
        "copy",
        "csv",
        "dataclasses",
        "datetime",
        "decimal",
        "enum",
        "fnmatch",
        "functools",
        "gc",
        "glob",
        "hashlib",
        "heapq",
        "hmac",
        "html",
        "http",
        "importlib",
        "inspect",
        "io",
        "itertools",
        "json",
        "logging",
        "math",
        "multiprocessing",
        "numbers",
        "operator",
        "os",
        "pathlib",
        "pickle",
        "platform",
        "pprint",
        "queue",
        "random",
        "re",
        "shutil",
        "signal",
        "socket",
        "sqlite3",
        "statistics",
        "string",
        "struct",
        "subprocess",
        "sys",
        "tempfile",
        "textwrap",
        "threading",
        "time",
        "traceback",
        "types",
        "typing",
        "unittest",
        "urllib",
        "uuid",
        "warnings",
        "weakref",
        "xml",
        "zipfile",
        "zlib",
        # common third-party
        "pytest",
        "numpy",
        "np",
        "pandas",
        "pd",
        "sklearn",
        "scipy",
        "torch",
        "yaml",
        "openpyxl",
        "requests",
        "httpx",
        "pydantic",
        "fastapi",
        "flask",
        "django",
        "matplotlib",
        "seaborn",
        "joblib",
        "tqdm",
        "click",
        "rich",
        "dotenv",
        "MetaTrader5",
        "mt5",
    }
)

_PROJECT_ROOTS = ("src", "tests", "scripts", "configs", "models", "multi_llm", "context")


def _humanize_token(name: str) -> str:
    """test_foo_bar / TestFooBar → readable phrase."""
    s = name
    if s.startswith("test_"):
        s = s[5:]
    elif s.startswith("Test"):
        s = s[4:]
    # CamelCase split
    s = re.sub(r"(?<!^)(?=[A-Z])", " ", s)
    s = s.replace("_", " ").strip()
    s = re.sub(r"\s+", " ", s)
    return s.lower() if s else name


def _doc_first_line(node: ast.AST | None) -> str:
    if node is None:
        return ""
    doc = ast.get_docstring(node) or ""
    if not doc:
        return ""
    line = doc.strip().splitlines()[0].strip()
    return re.sub(r"\s+", " ", line)


def _resolve_module_to_file(mod: str, file_path: Path) -> str | None:
    """Map import module string to a project-relative path if it looks local."""
    if not mod:
        return None
    root = mod.split(".", 1)[0]
    if root in _NON_PROJECT_ROOTS:
        return None
    # relative package under tests
    if root in _PROJECT_ROOTS or mod.startswith("."):
        # handle relative
        if mod.startswith("."):
            # count dots
            dots = len(mod) - len(mod.lstrip("."))
            remainder = mod.lstrip(".")
            base = file_path.parent
            for _ in range(dots - 1):
                base = base.parent
            if remainder:
                candidate = base / Path(*remainder.split("."))
            else:
                candidate = base
        else:
            candidate = REPO / Path(*mod.split("."))

        py = candidate.with_suffix(".py")
        init = candidate / "__init__.py"
        if py.is_file():
            return str(py.relative_to(REPO)).replace("\\", "/")
        if init.is_file():
            return str(init.relative_to(REPO)).replace("\\", "/")
        # best-effort path even if missing (still a referred target)
        rel = str(Path(*mod.split(".")).with_suffix(".py")).replace("\\", "/")
        if root in _PROJECT_ROOTS:
            return rel
        return rel

    # bare local import from same package (e.g. conftest helpers)
    sibling = file_path.parent / f"{root}.py"
    if sibling.is_file():
        return str(sibling.relative_to(REPO)).replace("\\", "/")
    # package under tests
    for base in (TESTS, REPO / "src", REPO / "scripts"):
        cand = base / root
        if (cand.with_suffix(".py")).is_file():
            return str(cand.with_suffix(".py").relative_to(REPO)).replace("\\", "/")
        if (cand / "__init__.py").is_file():
            return str((cand / "__init__.py").relative_to(REPO)).replace("\\", "/")
    return None


def extract_imports(tree: ast.AST, file_path: Path) -> list[str]:
    refs: list[str] = []
    seen: set[str] = set()

    def add(mod: str | None) -> None:
        if not mod:
            return
        resolved = _resolve_module_to_file(mod, file_path)
        if resolved and resolved not in seen:
            seen.add(resolved)
            refs.append(resolved)
        elif not resolved:
            # keep project-looking dotted names as-is if root is project-ish
            root = mod.lstrip(".").split(".", 1)[0]
            if root in _PROJECT_ROOTS and mod not in seen:
                seen.add(mod)
                refs.append(mod.replace(".", "/") + ".py")

    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                add(alias.name)
        elif isinstance(node, ast.ImportFrom):
            if node.module:
                base = (
                    ("." * node.level + node.module)
                    if node.level
                    else node.module
                )
                add(base)
                # also try package.submodule when `from pkg import submod`
                for alias in node.names:
                    if alias.name == "*":
                        continue
                    add(f"{base}.{alias.name}")
            elif node.level and node.level > 0:
                add("." * node.level)
                for alias in node.names:
                    if alias.name != "*":
                        add("." * node.level + alias.name)
    return sorted(refs)


def method_intent(fn: ast.FunctionDef | ast.AsyncFunctionDef) -> str:
    doc = _doc_first_line(fn)
    if doc:
        return doc
    return f"verifies {_humanize_token(fn.name)}"


def class_intent(cls: ast.ClassDef) -> str:
    """Combine all test methods' intents into one class intent."""
    class_doc = _doc_first_line(cls)
    methods: list[str] = []
    for item in cls.body:
        if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)):
            if item.name.startswith("test") or item.name in {
                "setup_method",
                "teardown_method",
                "setup_class",
                "teardown_class",
            }:
                if item.name.startswith("test"):
                    methods.append(method_intent(item))
    if not methods and not class_doc:
        return f"Test suite for {_humanize_token(cls.name)}"
    parts: list[str] = []
    if class_doc:
        parts.append(class_doc)
    if methods:
        # de-dupe while preserving order
        seen: set[str] = set()
        uniq: list[str] = []
        for m in methods:
            key = m.lower()
            if key not in seen:
                seen.add(key)
                uniq.append(m)
        parts.append("Methods cover: " + "; ".join(uniq))
    return " | ".join(parts)


def module_summary(
    tree: ast.Module,
    file_path: Path,
    classes: list[ast.ClassDef],
    free_tests: list[ast.FunctionDef | ast.AsyncFunctionDef],
) -> str:
    doc = _doc_first_line(tree)
    if doc:
        return doc
    # infer from filename
    stem = file_path.stem
    base = _humanize_token(stem if stem.startswith("test_") else f"test_{stem}")
    if classes:
        class_bits = ", ".join(c.name for c in classes[:6])
        more = f" (+{len(classes)-6} more)" if len(classes) > 6 else ""
        return (
            f"Pytest module covering {base}. "
            f"Defines test class(es): {class_bits}{more}."
        )
    if free_tests:
        return (
            f"Pytest module covering {base} via "
            f"{len(free_tests)} module-level test function(s)."
        )
    # helpers / fixtures / conftest / __init__
    if file_path.name == "conftest.py":
        return "Shared pytest fixtures and hooks for this package."
    if file_path.name == "__init__.py":
        return "Package marker / re-exports for the test package."
    if "fixture" in stem or "helper" in stem or "oracle" in stem:
        return f"Test support module: {base}."
    return f"Python test/support module for {base}."


def iter_py_files(root: Path) -> Iterable[Path]:
    for p in sorted(root.rglob("*.py")):
        if "__pycache__" in p.parts:
            continue
        yield p


def analyze_file(path: Path) -> list[dict[str, str]]:
    rel = str(path.relative_to(REPO)).replace("\\", "/")
    name = path.name
    try:
        src = path.read_text(encoding="utf-8", errors="replace")
        tree = ast.parse(src, filename=str(path))
    except SyntaxError as e:
        return [
            {
                "File Name": name,
                "Relative Path": rel,
                "Summary of functionality": f"Could not parse ({e})",
                "Referred files": "",
                "Test Class": "",
                "Intent of test class": "",
            }
        ]

    assert isinstance(tree, ast.Module)
    refs = extract_imports(tree, path)
    refs_str = "; ".join(refs) if refs else "(no project file imports)"

    classes: list[ast.ClassDef] = [
        n for n in tree.body if isinstance(n, ast.ClassDef)
    ]
    free_tests: list[ast.FunctionDef | ast.AsyncFunctionDef] = [
        n
        for n in tree.body
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
        and n.name.startswith("test")
    ]
    summary = module_summary(tree, path, classes, free_tests)

    rows: list[dict[str, str]] = []
    test_classes = [
        c
        for c in classes
        if c.name.startswith("Test")
        or any(
            isinstance(i, (ast.FunctionDef, ast.AsyncFunctionDef))
            and i.name.startswith("test")
            for i in c.body
        )
    ]
    # if no explicit Test* classes but has other classes with test methods
    if not test_classes and classes:
        # still report non-test helper classes? only if they have test_* methods
        test_classes = [
            c
            for c in classes
            if any(
                isinstance(i, (ast.FunctionDef, ast.AsyncFunctionDef))
                and i.name.startswith("test")
                for i in c.body
            )
        ]

    for cls in test_classes:
        rows.append(
            {
                "File Name": name,
                "Relative Path": rel,
                "Summary of functionality": summary,
                "Referred files": refs_str,
                "Test Class": cls.name,
                "Intent of test class": class_intent(cls),
            }
        )

    if free_tests:
        intents = [method_intent(f) for f in free_tests]
        seen: set[str] = set()
        uniq: list[str] = []
        for m in intents:
            k = m.lower()
            if k not in seen:
                seen.add(k)
                uniq.append(m)
        rows.append(
            {
                "File Name": name,
                "Relative Path": rel,
                "Summary of functionality": summary,
                "Referred files": refs_str,
                "Test Class": "(module-level tests)",
                "Intent of test class": "Methods cover: " + "; ".join(uniq),
            }
        )

    if not rows:
        # support / conftest / helpers without tests
        helper_classes = [c.name for c in classes]
        extra = (
            f" Classes: {', '.join(helper_classes)}." if helper_classes else ""
        )
        rows.append(
            {
                "File Name": name,
                "Relative Path": rel,
                "Summary of functionality": summary + extra,
                "Referred files": refs_str,
                "Test Class": "(no test class)",
                "Intent of test class": "Support/fixture module — no test methods defined.",
            }
        )
    return rows


def family_from_filename(name: str) -> str:
    """test_I_parent_htf_journeys.py → I; test_J_directional_displacement.py → J."""
    m = re.match(r"test_([A-Z])_", name)
    return m.group(1) if m else "?"


def family_from_grok_filename(name: str) -> str:   # back-compat alias
    return family_from_filename(name)


def _label_fields_from_doc(doc: str) -> dict[str, str]:
    """Line-oriented Source / Failure mode / Why / Numbers harvest."""
    fields = {lab: [] for lab in _GROK_LABELS}
    current: str | None = None
    for raw in doc.splitlines():
        line = raw.strip()
        if not line:
            continue
        hit = None
        for lab in _GROK_LABELS:
            if line.startswith(lab):
                hit = lab
                break
        if hit:
            current = hit
            rest = line[len(hit) :].lstrip()
            if rest:
                fields[hit].append(rest)
            continue
        if current:
            fields[current].append(line)
    return {
        lab: re.sub(r"\s+", " ", " ".join(parts)).strip()
        for lab, parts in fields.items()
    }


def parse_grok_docstring(doc: str | None, fn_name: str) -> dict[str, str]:
    """Extract Intent / Trying-to-do / labeled fields per design §4."""
    inferred = f"verifies {_humanize_token(fn_name)}"
    if not doc or not doc.strip():
        return {
            "Intent": inferred,
            "Trying to do": inferred,
            "Source contract": "",
            "Failure mode": "",
            "Why ordinary tests miss it": "",
            "Intent quality": "INTENT_INFERRED",
        }

    labels = _label_fields_from_doc(doc)
    unlabeled: list[str] = []
    buf: list[str] = []
    seen_label = False

    def flush() -> None:
        nonlocal buf
        if buf:
            unlabeled.append(re.sub(r"\s+", " ", " ".join(buf)).strip())
            buf = []

    for raw in doc.strip().splitlines():
        stripped = raw.strip()
        if not stripped:
            flush()
            continue
        if any(stripped.startswith(lab) for lab in _GROK_LABELS):
            seen_label = True
            flush()
            continue
        if seen_label:
            # continuation of a labeled field — already harvested
            continue
        buf.append(stripped)
    flush()

    intent = unlabeled[0] if unlabeled else ""
    trying = " ".join(unlabeled[1:]).strip()
    source = labels["Source:"]
    fail = labels["Failure mode:"]
    why = labels["Why ordinary tests miss it:"]

    if not intent:
        intent = inferred
        quality = "INTENT_PARTIAL"
    elif source and fail:
        quality = "FROM_DOCSTRING"
    else:
        quality = "INTENT_PARTIAL"
    if not trying:
        trying = intent
    return {
        "Intent": intent,
        "Trying to do": trying,
        "Source contract": source,
        "Failure mode": fail,
        "Why ordinary tests miss it": why,
        "Intent quality": quality,
    }


def analyze_suite_methods(path: Path) -> list[dict]:
    """AST: each def test_* plus docstring fields and referred files."""
    src = path.read_text(encoding="utf-8", errors="replace")
    tree = ast.parse(src, filename=str(path))
    refs = extract_imports(tree, path)
    refs_str = "; ".join(refs) if refs else "(no project file imports)"
    rel = str(path.relative_to(REPO)).replace("\\", "/")
    family = family_from_filename(path.name)
    layer = _FAMILY_LAYER.get(family, "UNMAPPED")
    out: list[dict] = []
    for node in tree.body:
        if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            continue
        if not node.name.startswith("test"):
            continue
        fields = parse_grok_docstring(ast.get_docstring(node), node.name)
        out.append(
            {
                "Family": family,
                "Topic layer": layer,
                "File": rel,
                "Method": node.name,
                "Referred files": refs_str,
                **fields,
            }
        )
    return out


def analyze_grok_methods(path: Path) -> list[dict]:   # back-compat alias
    return analyze_suite_methods(path)


def collect_nodeids(suite: str | Suite = "Grok") -> list[str]:
    """Collector contract: sys.executable -m pytest --collect-only -q tests/<Suite>."""
    sui = _suite(suite)
    rel_dir = str(sui.directory.relative_to(REPO)).replace("\\", "/")
    try:
        proc = subprocess.run(
            [sys.executable, "-m", "pytest", "--collect-only", "-q", rel_dir],
            cwd=REPO,
            timeout=120,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
    except subprocess.TimeoutExpired as exc:
        raise CollectFailed(f"collect timed out after 120s: {exc}") from exc
    except OSError as exc:
        raise CollectFailed(f"collect could not start: {exc}") from exc
    if proc.returncode not in (0,):
        tail = (proc.stderr or proc.stdout or "")[-800:]
        raise CollectFailed(f"collect exit {proc.returncode}: {tail}")
    nodeids: list[str] = []
    for line in (proc.stdout or "").splitlines():
        s = line.strip().replace("\\", "/")
        if "::" not in s:
            continue
        prefix = f"{rel_dir}/"
        if s.startswith(prefix) or f"/{prefix}" in s:
            # keep repo-relative form
            if not s.startswith(prefix):
                s = s[s.find(prefix):]
            nodeids.append(s)
    if not nodeids:
        raise CollectFailed("collect produced 0 nodeids")
    return nodeids


def collect_grok_nodeids() -> list[str]:   # back-compat alias
    return collect_nodeids("Grok")


def _case_id_from_nodeid(nodeid: str) -> str:
    if "[" in nodeid and nodeid.endswith("]"):
        return nodeid[nodeid.rfind("[") + 1 : -1]
    return "-"


def _method_from_nodeid(nodeid: str) -> str:
    tail = nodeid.split("::", 1)[-1]
    return tail.split("[", 1)[0]


def _file_from_nodeid(nodeid: str) -> str:
    return nodeid.split("::", 1)[0].replace("\\", "/")


def build_intent_rows(
    suite: str | Suite = "Grok", nodeids: list[str] | None = None
) -> list[dict]:
    """Join collected nodeids to AST function metadata."""
    sui = _suite(suite)
    if nodeids is None:
        nodeids = collect_nodeids(sui)
    by_file_method: dict[tuple[str, str], dict] = {}
    arity: dict[tuple[str, str], int] = {}
    for nid in nodeids:
        key = (_file_from_nodeid(nid), _method_from_nodeid(nid))
        arity[key] = arity.get(key, 0) + 1
    for path in sorted(sui.directory.glob("test_*.py")):
        for rec in analyze_suite_methods(path):
            by_file_method[(rec["File"], rec["Method"])] = rec
    rows: list[dict] = []
    for nid in nodeids:
        rel = _file_from_nodeid(nid)
        method = _method_from_nodeid(nid)
        meta = by_file_method.get((rel, method))
        if meta is None:
            family = family_from_filename(Path(rel).name)
            meta = {
                "Family": family,
                "Topic layer": _FAMILY_LAYER.get(family, "UNMAPPED"),
                "File": rel,
                "Method": method,
                "Intent": f"verifies {_humanize_token(method)}",
                "Trying to do": f"verifies {_humanize_token(method)}",
                "Source contract": "",
                "Failure mode": "",
                "Why ordinary tests miss it": "",
                "Intent quality": "INTENT_INFERRED",
                "Referred files": "",
            }
        rows.append(
            {
                **meta,
                "Nodeid": nid,
                "Case id": _case_id_from_nodeid(nid),
                "Parametrize arity": arity.get((rel, method), 1),
                "Added by": sui.name,
            }
        )
    return rows


def build_grok_intent_rows(nodeids: list[str] | None = None) -> list[dict]:
    """Back-compat alias — the Grok suite at nodeid grain."""
    return build_intent_rows("Grok", nodeids)


def _style_header_row(ws, headers: list[str]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_intent_excel(
    rows: list[dict], out: Path, suite: str | Suite = "Grok"
) -> None:
    """Sibling workbook: <Suite> Method Intent + By Method + README."""
    sui = _suite(suite)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sui.sheet
    headers = _INTENT_HEADERS
    _style_header_row(ws, headers)
    wrap = Alignment(wrap_text=True, vertical="top")
    for r_i, row in enumerate(rows, 2):
        for c_i, h in enumerate(headers, 1):
            val = row.get(h, "")
            cell = ws.cell(r_i, c_i, val)
            cell.alignment = wrap
    widths = {
        "A": 10, "B": 36, "C": 44, "D": 48, "E": 70,
        "F": 28, "G": 55, "H": 55, "I": 48, "J": 48,
        "K": 40, "L": 18, "M": 16, "N": 12, "O": 50,
    }
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w
    last = max(len(rows) + 1, 2)
    ws.auto_filter.ref = f"A1:O{last}"
    ws.freeze_panes = "A2"

    # By Method rollup
    ws2 = wb.create_sheet("By Method")
    by_headers = [h for h in headers]
    _style_header_row(ws2, by_headers)
    grouped: dict[tuple[str, str], list[dict]] = {}
    for row in rows:
        grouped.setdefault((row["File"], row["Method"]), []).append(row)
    r = 2
    for (_file, _method), group in grouped.items():
        first = group[0]
        case_ids = [g["Case id"] for g in group]
        case_cell = (
            "(unparametrized)"
            if case_ids == ["-"]
            else "; ".join(case_ids)
        )
        values = {
            **first,
            "Nodeid": "",
            "Case id": case_cell,
            "Parametrize arity": len(group),
        }
        for c_i, h in enumerate(by_headers, 1):
            cell = ws2.cell(r, c_i, values.get(h, ""))
            cell.alignment = wrap
        r += 1
    for letter, w in widths.items():
        ws2.column_dimensions[letter].width = w
    ws2.auto_filter.ref = f"A1:O{max(r - 1, 2)}"
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("README")
    rel_dir = str(sui.directory.relative_to(REPO)).replace("\\", "/")
    notes = [
        f"{sui.name} auditor suite — one row per pytest nodeid (test case)",
        "Generated by scripts/analysis/test_functionality_excel.py (SCR-359)",
        "Regenerate: python scripts/analysis/test_functionality_excel.py",
        f"Intent is extracted from {rel_dir} docstrings. Do not hand-edit this book.",
        "INTENT_INFERRED = no docstring. INTENT_PARTIAL = docstring missing Source: or Failure mode:.",
        f"Nodeid rows: {len(rows)}",
        f"Function rows: {len(grouped)}",
    ]
    for i, line in enumerate(notes, 1):
        ws3.cell(i, 1, line)
    ws3.column_dimensions["A"].width = 120
    wb.save(out)


def write_grok_intent_excel(rows: list[dict], out: Path) -> None:
    """Back-compat alias — writes the Grok-titled sibling."""
    write_intent_excel(rows, out, "Grok")


def write_collect_failed_stub(out: Path, detail: str, suite: str | Suite = "Grok") -> None:
    """Overwrite sibling so a stale good run cannot sit next to a failed collect."""
    from datetime import datetime, timezone

    sui = _suite(suite)
    stub = {
        "Family": "-",
        "Topic layer": "UNMAPPED",
        "File": str(sui.directory.relative_to(REPO)).replace("\\", "/"),
        "Method": "-",
        "Nodeid": "",
        "Case id": "-",
        "Intent": "COLLECT_FAILED",
        "Trying to do": "COLLECT_FAILED",
        "Source contract": "",
        "Failure mode": detail[:500],
        "Why ordinary tests miss it": "",
        "Intent quality": "INTENT_INFERRED",
        "Parametrize arity": 0,
        "Added by": sui.name,
        "Referred files": "",
    }
    write_intent_excel([stub], out, sui)
    # stamp README with UTC + detail
    wb = openpyxl_load(out)
    ws = wb["README"]
    ws.cell(8, 1, f"COLLECT_FAILED at {datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')}")
    ws.cell(9, 1, detail[:1500])
    wb.save(out)


def write_grok_collect_failed_stub(out: Path, detail: str) -> None:
    """Back-compat alias — Grok-titled COLLECT_FAILED stub."""
    write_collect_failed_stub(out, detail, "Grok")


def openpyxl_load(path: Path):
    from openpyxl import load_workbook

    return load_workbook(path)


def write_excel(
    rows: list[dict[str, str]],
    out: Path,
    grok_pointer: dict[str, str] | None = None,
    claude_pointer: dict[str, str] | None = None,
) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Functionality"

    headers = [
        "File Name",
        "Relative Path",
        "Summary of functionality",
        "Referred files",
        "Test Class",
        "Intent of test class",
    ]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for r_i, row in enumerate(rows, 2):
        for c_i, h in enumerate(headers, 1):
            cell = ws.cell(r_i, c_i, row.get(h, ""))
            cell.alignment = wrap

    widths = {
        "A": 42,
        "B": 48,
        "C": 55,
        "D": 55,
        "E": 36,
        "F": 70,
    }
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    # second sheet: one row per file (classes collapsed)
    ws2 = wb.create_sheet("By File")
    for col, h in enumerate(
        ["File Name", "Relative Path", "Summary of functionality", "Referred files", "Test Classes", "Intent of test class"],
        1,
    ):
        cell = ws2.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    by_file: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        by_file.setdefault(row["Relative Path"], []).append(row)

    r = 2
    for rel, group in sorted(by_file.items()):
        first = group[0]
        classes = " | ".join(
            g["Test Class"] for g in group if g["Test Class"]
        )
        intents = " || ".join(
            f"[{g['Test Class']}] {g['Intent of test class']}" for g in group
        )
        values = [
            first["File Name"],
            rel,
            first["Summary of functionality"],
            first["Referred files"],
            classes,
            intents,
        ]
        for c, v in enumerate(values, 1):
            cell = ws2.cell(r, c, v)
            cell.alignment = wrap
        r += 1

    for letter, w in {"A": 42, "B": 48, "C": 55, "D": 55, "E": 40, "F": 80}.items():
        ws2.column_dimensions[letter].width = w
    ws2.auto_filter.ref = f"A1:F{r-1}"
    ws2.freeze_panes = "A2"

    # readme sheet
    ws3 = wb.create_sheet("README")
    notes = [
        "tests/ Python inventory — business functionality (direct imports only)",
        f"Generated from: {TESTS}",
        "Scope: every .py under tests/ (excluding __pycache__)",
        "Summary of functionality: module docstring first line, else inferred from file/class names",
        "Referred files: direct import targets resolved to project-relative paths (stdlib/third-party omitted)",
        "Intent of test class: class docstring + combined intents of all test_* methods (docstring or humanized name)",
        "Sheet 'Test Functionality': one row per test class (or module-level tests / support module)",
        "Sheet 'By File': one row per file with classes collapsed",
        "No nested/transitive dependency analysis — only imports written in each file",
    ]
    for i, line in enumerate(notes, 1):
        ws3.cell(i, 1, line)
    ws3.column_dimensions["A"].width = 120

    ptr_headers = [
        "Sibling path",
        "Status",
        "Generated-at UTC",
        "Family count",
        "Function count",
        "Nodeid count",
        "Note",
    ]
    # One pointer sheet per auditor suite. These are NEW sheets, never new columns
    # on "Test Functionality" / "By File" — the A-F positional contract pinned by
    # tests/test_semantic_identity_workbooks.py stays untouched.
    for suite_name, pointer in (("Grok", grok_pointer), ("Claude", claude_pointer)):
        sui = _SUITES[suite_name]
        ws_ptr = wb.create_sheet(sui.pointer_sheet)
        _style_header_row(ws_ptr, ptr_headers)
        ptr = pointer or {
            "Sibling path": str(sui.out.relative_to(REPO)).replace("\\", "/"),
            "Status": "NOT_RUN",
            "Generated-at UTC": "",
            "Family count": "",
            "Function count": "",
            "Nodeid count": "",
            "Note": f"{sui.name} intent not generated in this run",
        }
        for c, h in enumerate(ptr_headers, 1):
            ws_ptr.cell(2, c, ptr.get(h, ""))
        for letter, w in {
            "A": 42, "B": 16, "C": 22, "D": 14, "E": 16, "F": 14, "G": 70,
        }.items():
            ws_ptr.column_dimensions[letter].width = w
        ws_ptr.freeze_panes = "A2"

    wb.save(out)


def _utc_now() -> str:
    from datetime import datetime, timezone

    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _pointer(
    status: str,
    families: int | str = "",
    functions: int | str = "",
    nodeids: int | str = "",
    note: str = "",
    suite: str | Suite = "Grok",
) -> dict[str, str]:
    return {
        "Sibling path": str(_suite(suite).out.relative_to(REPO)).replace("\\", "/"),
        "Status": status,
        "Generated-at UTC": _utc_now(),
        "Family count": str(families),
        "Function count": str(functions),
        "Nodeid count": str(nodeids),
        "Note": note,
    }


def emit_intent(suite: str | Suite = "Grok", *, require: bool) -> tuple[dict[str, str], int]:
    """Write one suite's sibling workbook. Returns (pointer, exit_extra)."""
    sui = _suite(suite)
    rel_out = str(sui.out.relative_to(REPO)).replace("\\", "/")
    tag = f"{sui.name.lower()}_intent"
    try:
        nodeids = collect_nodeids(sui)
        rows = build_intent_rows(sui, nodeids)
        write_intent_excel(rows, sui.out, sui)
        families = sorted({r["Family"] for r in rows if r["Family"] not in {"-", "?"}})
        functions = {(r["File"], r["Method"]) for r in rows}
        ptr = _pointer(
            "OK",
            families=len(families),
            functions=len(functions),
            nodeids=len(rows),
            note=rel_out,
            suite=sui,
        )
        print(f"{tag} nodeids={len(rows)} functions={len(functions)} out={sui.out}")
        return ptr, 0
    except CollectFailed as exc:
        write_collect_failed_stub(sui.out, str(exc), sui)
        ptr = _pointer("COLLECT_FAILED", note=str(exc)[:300], suite=sui)
        print(f"{tag} COLLECT_FAILED: {exc}", file=sys.stderr)
        return ptr, (1 if require else 0)


def emit_grok_intent(*, require: bool) -> tuple[dict[str, str], int]:
    """Back-compat alias."""
    return emit_intent("Grok", require=require)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Inventory tests/**/*.py and emit the Grok + Claude test-intent Excels. "
            "Default: class-grain book always, both siblings best-effort. "
            "--grok-only and --claude-only are separate groups, so passing both means "
            "'both siblings, no class-grain book'."
        )
    )
    grok = parser.add_mutually_exclusive_group()
    grok.add_argument(
        "--grok-only",
        action="store_true",
        help="Write only the sibling grok_test_intent.xlsx (non-zero if collect fails).",
    )
    grok.add_argument(
        "--skip-grok-intent",
        action="store_true",
        help="Do not invoke pytest for tests/Grok. Do not touch the Grok sibling.",
    )
    grok.add_argument(
        "--require-grok-intent",
        action="store_true",
        help="Non-zero exit if the tests/Grok collect fails.",
    )
    claude = parser.add_mutually_exclusive_group()
    claude.add_argument(
        "--claude-only",
        action="store_true",
        help="Write only the sibling claude_test_intent.xlsx (non-zero if collect fails).",
    )
    claude.add_argument(
        "--skip-claude-intent",
        action="store_true",
        help="Do not invoke pytest for tests/Claude. Do not touch the Claude sibling.",
    )
    claude.add_argument(
        "--require-claude-intent",
        action="store_true",
        help="Non-zero exit if the tests/Claude collect fails.",
    )
    args = parser.parse_args(argv)

    # Sibling-only mode: write the named sibling(s) and nothing else.
    if args.grok_only or args.claude_only:
        extra = 0
        if args.grok_only:
            _ptr, code = emit_intent("Grok", require=True)
            extra |= code
        if args.claude_only:
            _ptr, code = emit_intent("Claude", require=True)
            extra |= code
        return extra

    rows: list[dict[str, str]] = []
    files = list(iter_py_files(TESTS))
    for p in files:
        rows.extend(analyze_file(p))

    extra = 0
    if args.skip_grok_intent:
        grok_ptr = _pointer(
            "SKIPPED", note="--skip-grok-intent; sibling not touched", suite="Grok"
        )
    else:
        grok_ptr, code = emit_intent("Grok", require=args.require_grok_intent)
        extra |= code

    if args.skip_claude_intent:
        claude_ptr = _pointer(
            "SKIPPED", note="--skip-claude-intent; sibling not touched", suite="Claude"
        )
    else:
        claude_ptr, code = emit_intent("Claude", require=args.require_claude_intent)
        extra |= code

    write_excel(rows, OUT, grok_pointer=grok_ptr, claude_pointer=claude_ptr)
    print(f"files={len(files)} rows={len(rows)} out={OUT}")
    return extra


if __name__ == "__main__":
    sys.exit(main())
