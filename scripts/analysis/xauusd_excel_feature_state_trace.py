#!/usr/bin/env python3
"""
xauusd_excel_feature_state_trace.py
====================================
Observe-only per-bar trace over an MT5-exported XAUUSD workbook:

    Excel/CSV OHLCV -> FeaturePipeline (39 canonical features) -> CRTEngine states

Emits ONE ROW PER BAR joining the raw candle, the full canonical feature vector, and the
CRT state transition (state_before -> action -> state_after), plus a read-only 4-engine
fusion overlay so you can see which committed entries the gate would veto.

Read-only by construction: no src/ change, no config write, no new formula or tunable. It
renders math and state that already exist.

Two deliberate design choices, both about not lying to the reader
-----------------------------------------------------------------
1. WARMUP IS SHOWN, NOT FILLED. FeaturePipeline.run() ends with finalize(), which drops every
   row where any canonical feature is still NaN (rolling warmup). On a ~1.2k-bar corpus that
   would delete the head of the very window being traced. So this script calls the SAME public
   compute_*/promote_* methods run() calls, in the SAME order, and deliberately never calls
   finalize(). Warmup cells are left EMPTY (NaN) and flagged by `warmup_complete` -- never
   ffill'd, bfill'd, or zero-filled into a fabricated number.

   Note the CRT engine keeps its OWN internal ATR (EngineState.atr_abs) independent of the
   feature pipeline, so the two tracks are parallel, not chained: warmup NaNs in the feature
   columns do not affect the state machine.

2. THE GATE OVERLAY IS AN OVERLAY, NOT A SECOND REPLAY. Pass A is the CRT trajectory
   (BACKTEST_ENGINE_GATE=0 semantics -- the F-037 CRT-isolated research spine). Pass B runs
   EngineRunner -> fusion -> decision over the same bars WITHOUT feeding any veto back into
   the state machine, so both passes share one trajectory. A real gate-ON replay diverges:
   F-055 showed a reject RESETS the CRT state machine, so removed != added. These columns
   answer "which committed entries would fusion veto", NOT "what would the ledger look like".
   F-070 measured 0/30 vetoes on the active config, so near-identical columns are the expected
   result here, not a bug.

Usage:
    venv/Scripts/python.exe scripts/analysis/xauusd_excel_feature_state_trace.py
    venv/Scripts/python.exe scripts/analysis/xauusd_excel_feature_state_trace.py \\
        --csv data/XAUUSD_M15.csv --overlay events
    venv/Scripts/python.exe scripts/analysis/xauusd_excel_feature_state_trace.py \\
        --xlsx data/XAUUSD_M15_20260807_203705.xlsx --output-dir results/xauusd_excel_trace
    venv/Scripts/python.exe scripts/analysis/xauusd_excel_feature_state_trace.py \\
        --xlsx data/XAUUSD_M15_20260807_200204.xlsx \\
        --htf-bars 16 --tp1-r 1.5 --tp2-r 3.0 --overlay none \\
        --output-dir results/xauusd_excel_trace

Knobs (--htf-bars, --tp1-r, --tp2-r, --sl-atr-buffer, --shadow-ttl,
--ignore-off-session) are THIS-RUN overrides. They do not write production JSON.
Each is listed in knobs.json.

--research-xy sweeps x=tp1, y=tp2, z=sl_atr_buffer on RETEST_CONFIRMED after
the CRT replay (geometry only). Feature-pipeline config + that bar's feature
values are stamped per setup.
"""
from __future__ import annotations

import argparse
import hashlib
import inspect
import json
import logging
import os
import re
import sys
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
os.chdir(ROOT)

# CRT-isolated trajectory (F-037): the state machine is replayed without a fusion veto.
# Pass B below is a read-only overlay on top of that same trajectory.
os.environ.setdefault("BACKTEST_ENGINE_GATE", "0")

import numpy as np
import pandas as pd

from config_layer.crt_engine_v2 import CRTEngine, Candle
from config_layer.production_config import (
    get_active_version,
    get_prod_metadata,
    get_prod_section,
    load_prod_config_from_registry,
)
from core.engine_runner import EngineRunner
from features.feature_pipeline import FeaturePipeline, _FP_CFG_KEYS, required_warmup_rows
from features.feature_schema import CANONICAL_FEATURES, SCHEMA_VERSION
from runtime.backtest_v2 import HTFBuilder

logger = logging.getLogger("XAUUSD_EXCEL_TRACE")
logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")

DEFAULT_XLSX = ROOT / "data" / "XAUUSD_M15_20260807_203705.xlsx"
DEFAULT_OUTPUT_DIR = ROOT / "results" / "crt_survey_trace"
DEFAULT_INSTRUMENT = "XAUUSD"


def corpus_for(instrument: str) -> Path:
    """Default corpus for an instrument: the canonical MT5 M15 export."""
    return ROOT / "data" / "mt5" / f"{instrument}_M15.csv"

TIMEFRAME = "M15"

# ─────────────────────────────────────────────────────────────────────────────
# PIPELINE-STAGE DRIFT GUARD
#
# Verbatim copy of the ordered method-call sequence inside FeaturePipeline.run(), EXCLUDING
# finalize()/log_critical_feature_health()/build_feature_vector(). If run() is ever edited,
# _assert_stage_list_matches_run_source() fails LOUDLY at startup rather than silently tracing
# a stale sequence. (Same guard idea as scripts/analysis/feature_trace_report.py.)
# ─────────────────────────────────────────────────────────────────────────────
PIPELINE_STAGES = [
    "compute_price_features", "compute_volume_features", "compute_indicators",
    "compute_trend_features", "compute_volatility_regime", "compute_context",
    "compute_structure_liquidity", "compute_normalization",
    "compute_canonical_price_features", "compute_canonical_volatility_features",
    "compute_canonical_ema_features", "compute_canonical_trend_features",
    "compute_canonical_structure_features", "compute_canonical_temporal_features",
    "compute_liquidity_distance", "promote_volume_spike", "compute_canonical_session",
]
_STAGES_EXCLUDED_FROM_TRACE = {
    "finalize", "log_critical_feature_health", "build_feature_vector",
}

# CRT actions that terminate or gate a setup — always interesting in the trace.
TERMINAL_ACTIONS = {
    "DISPLACEMENT_CONFIRMED", "EXPANSION_CONFIRMED", "RETEST_CONFIRMED", "TRADE_OPENED",
    "FILTER_REJECTED", "SWEEP_EXPIRED", "EXPANSION_EXPIRED", "SHADOW_ADVISORY_BLOCK",
    "CONFIRMATION_FAILED",
}


def _assert_stage_list_matches_run_source() -> None:
    """Fail loudly if FeaturePipeline.run()'s stage order drifts from PIPELINE_STAGES."""
    src = inspect.getsource(FeaturePipeline.run)
    called = [
        m for m in re.findall(r"self\.(\w+)\(\)", src)
        if m not in _STAGES_EXCLUDED_FROM_TRACE
    ]
    if called != PIPELINE_STAGES:
        raise RuntimeError(
            "FeaturePipeline.run() stage sequence drifted from this script's PIPELINE_STAGES.\n"
            f"  run()  : {called}\n"
            f"  script : {PIPELINE_STAGES}\n"
            "Update PIPELINE_STAGES to match run() before trusting this trace."
        )


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ─────────────────────────────────────────────────────────────────────────────
# 1. LOAD
# ─────────────────────────────────────────────────────────────────────────────
def _to_candles(rows: list[tuple]) -> list[Candle]:
    candles: list[Candle] = []
    for i, (ts, o, h, l, c, v) in enumerate(rows):
        if isinstance(ts, str):
            ts = datetime.fromisoformat(ts)
        candles.append(
            Candle(
                timestamp=ts,
                open=float(o), high=float(h), low=float(l), close=float(c),
                volume=float(v or 0),
                index=i,
            )
        )
    return candles


def load_from_xlsx(path: Path) -> tuple[list[Candle], dict[str, Any]]:
    """Read the `candles` sheet -> list[Candle]; carry `source_info` provenance if present.

    Loader shape lifted from reports/xauusd_retest_pathb_counterfactual_probe.py:74-94.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb["candles"].iter_rows(min_row=2, values_only=True))
        source_info: dict[str, Any] = {}
        if "source_info" in wb.sheetnames:
            sheet = wb["source_info"]
            data = list(sheet.iter_rows(values_only=True))
            if len(data) >= 2:
                headers = [str(h) for h in data[0]]
                values = data[1]
                # `login` is an account identifier — deliberately never carried into output.
                source_info = {
                    k: v for k, v in zip(headers, values) if k.lower() != "login"
                }
    finally:
        wb.close()
    return _to_candles(rows), source_info


def load_from_csv(path: Path) -> tuple[list[Candle], dict[str, Any]]:
    df = pd.read_csv(path)
    rows = list(
        df[["timestamp", "open", "high", "low", "close", "volume"]].itertuples(
            index=False, name=None
        )
    )
    return _to_candles(rows), {}


# ─────────────────────────────────────────────────────────────────────────────
# 2. FEATURES  (all stages, no finalize -> every bar survives, warmup stays NaN)
# ─────────────────────────────────────────────────────────────────────────────
def _snapshot_feature_cfg(pipe: FeaturePipeline) -> dict[str, Any]:
    """Keys the live FeaturePipeline instance actually resolved — not a comment dump."""
    cfg = pipe._fp_cfg
    snap: dict[str, Any] = {}
    for key in _FP_CFG_KEYS:
        snap[key] = cfg[key]
    return snap


def build_feature_frame(candles: list[Candle]) -> tuple[pd.DataFrame, dict[str, Any]]:
    df = pd.DataFrame(
        {
            "timestamp": [c.timestamp for c in candles],
            "open":   [c.open for c in candles],
            "high":   [c.high for c in candles],
            "low":    [c.low for c in candles],
            "close":  [c.close for c in candles],
            "volume": [c.volume for c in candles],
        }
    )
    pipe = FeaturePipeline(df)
    for stage in PIPELINE_STAGES:
        getattr(pipe, stage)()

    out = pipe.df.replace([np.inf, -np.inf], np.nan).copy()
    for col in CANONICAL_FEATURES:
        if col not in out.columns:
            # Absent slot is a real defect, not a warmup gap — surface it, don't fill it.
            logger.warning("canonical feature absent from pipeline output: %s", col)
            out[col] = np.nan
    if len(out) != len(candles):
        raise RuntimeError(
            f"feature frame length {len(out)} != candle count {len(candles)}; "
            "a stage dropped rows (finalize must not run here)."
        )
    out["warmup_complete"] = out[list(CANONICAL_FEATURES)].notna().all(axis=1)
    return out, _snapshot_feature_cfg(pipe)


# ─────────────────────────────────────────────────────────────────────────────
# 3. STATES  (CRT replay — production sessions unless --ignore-off-session)
# ─────────────────────────────────────────────────────────────────────────────
def replay_crt(
    candles: list[Candle],
    instrument: str,
    *,
    htf_bars: int = 4,
    crt_cfg: Any = None,
) -> tuple[list[dict[str, Any]], Any]:
    """Replay CRTEngine bar by bar; return one record per bar (index-aligned)."""
    if crt_cfg is None:
        crt_cfg = load_prod_config_from_registry(get_active_version(), instrument)
    engine = CRTEngine(crt_cfg)
    htf = HTFBuilder(htf_bars, instrument)
    initialized = False

    # Capture ev_log events and attach them to the bar they fired on.
    per_bar_events: dict[int, list[dict[str, Any]]] = {}
    current_index = {"i": -1}
    orig_record = engine.ev_log.record

    def record(event, candle, **kwargs):  # type: ignore[no-untyped-def]
        per_bar_events.setdefault(current_index["i"], []).append(
            {
                "event": str(getattr(event, "name", event)),
                "reason": kwargs.get("reason"),
                "direction": _plain(kwargs.get("direction")),
                "price": kwargs.get("price"),
                "metadata": _plain(kwargs.get("metadata")),
            }
        )
        return orig_record(event, candle, **kwargs)

    engine.ev_log.record = record  # type: ignore[method-assign]

    records: list[dict[str, Any]] = []
    for candle in candles:
        current_index["i"] = candle.index
        completed = htf.push(candle)

        if not initialized:
            if completed:
                engine.initialise_range(htf.seed_candles(), htf.current_htf_id, "UNKNOWN")
                initialized = True
            records.append(_bar_record(candle, None, None, None, engine, per_bar_events, seeded=False))
            continue

        before = engine.state.current_state.name
        result = engine.process_candle(candle, htf.current_htf_id)
        after = engine.state.current_state.name
        action = result.get("action") if isinstance(result, dict) else str(result)
        records.append(
            _bar_record(candle, before, after, action, engine, per_bar_events, seeded=True)
        )

    engine.ev_log.record = orig_record  # type: ignore[method-assign]
    return records, engine


def _bar_record(
    candle: Candle,
    before: str | None,
    after: str | None,
    action: str | None,
    engine: CRTEngine,
    per_bar_events: dict[int, list[dict[str, Any]]],
    *,
    seeded: bool,
) -> dict[str, Any]:
    st = engine.state
    rng = st.active_range
    sw = st.sweep_event
    disp = st.displacement_candle
    rt = st.retest_candle
    rs = st.risk_score
    trade = st.active_trade
    events = per_bar_events.get(candle.index, [])

    rec: dict[str, Any] = {
        "bar_index": candle.index,
        "timestamp": candle.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
        "htf_seeded": seeded,
        "state_before": before,
        "state_after": after,
        "action": action,
        "direction": st.direction.name if st.direction is not None else None,
        "risk_score": float(rs.final) if rs is not None else None,
        "crt_atr_abs": float(st.atr_abs),
        "ema_fast_val": float(st.ema_fast_val),
        "ema_slow_val": float(st.ema_slow_val),
        "range_session": rng.session if rng else None,
        "event": ";".join(e["event"] for e in events) or None,
        "reject_reason": ";".join(
            str(e["reason"]) for e in events if e.get("reason")
        ) or None,
        # nested block for the JSONL side (shape follows crt_xauusd_runtime_trace.py:195-260)
        "_crt_inputs": {
            "active_range": (
                {
                    "h_ref": rng.h_ref, "l_ref": rng.l_ref,
                    "equilibrium": rng.equilibrium, "size": rng.size,
                    "htf_candle_id": rng.htf_candle_id, "session": rng.session,
                }
                if rng else None
            ),
            "sweep": (
                {
                    "direction": _plain(sw.direction), "price": sw.price,
                    "double_confirmed": sw.double_confirmed,
                    "candle_index": sw.candle_index, "sweep_type": sw.sweep_type,
                }
                if sw else None
            ),
            "displacement_ohlc": (
                {"open": disp.open, "high": disp.high, "low": disp.low,
                 "close": disp.close, "index": disp.index} if disp else None
            ),
            "retest_ohlc": (
                {"open": rt.open, "high": rt.high, "low": rt.low,
                 "close": rt.close, "index": rt.index} if rt else None
            ),
            "cached_features": dict(st.cached_features) if st.cached_features else None,
            "events": events,
        },
    }

    if trade is not None and action == "TRADE_OPENED":
        rec["trade_entry"] = float(trade.entry_price)
        rec["trade_sl"] = float(trade.sl_price)
        rec["trade_tp1"] = float(trade.tp1_price)
        rec["trade_tp2"] = float(trade.tp2_price)
        rec["trade_risk_pct"] = float(trade.risk_pct)
        rec["trade_id"] = trade.id
    return rec


def _plain(obj: Any) -> Any:
    """Enum/dataclass -> JSON-safe scalar."""
    if obj is None or isinstance(obj, (str, int, float, bool)):
        return obj
    if isinstance(obj, dict):
        return {str(k): _plain(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_plain(v) for v in obj]
    if hasattr(obj, "value") and not callable(obj.value):
        return obj.value
    if hasattr(obj, "name"):
        return obj.name
    try:
        return float(obj)
    except Exception:
        return str(obj)


def _rel(path: Path) -> str:
    path = Path(path).resolve()
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def planned_entry_sl_tp(
    rec: dict[str, Any],
    *,
    tp1_r: float,
    tp2_r: float,
    sl_atr_buffer: float,
) -> dict[str, Any] | None:
    """Same geometry as ExecutionEngine.build_trade, using this run's R knobs.

    Computed on RETEST_CONFIRMED even if session later rejects — so the trace
    still shows Entry/SL/TP when the engine never opens.
    """
    crt = rec.get("_crt_inputs") or {}
    disp = crt.get("displacement_ohlc")
    rt = crt.get("retest_ohlc")
    direction = rec.get("direction")
    atr = rec.get("crt_atr_abs")
    if not disp or not rt or direction not in ("LONG", "SHORT") or atr is None:
        return None
    entry = float(rt["close"])
    atr_f = float(atr)
    if direction == "LONG":
        sl = float(disp["low"]) - sl_atr_buffer * atr_f
        sl_valid = sl < entry
    else:
        sl = float(disp["high"]) + sl_atr_buffer * atr_f
        sl_valid = sl > entry
    risk = abs(entry - sl)
    if direction == "LONG":
        tp1 = entry + tp1_r * risk
        tp2 = entry + tp2_r * risk
    else:
        tp1 = entry - tp1_r * risk
        tp2 = entry - tp2_r * risk
    return {
        "plan_direction": direction,
        "plan_entry": entry,
        "plan_sl": sl,
        "plan_tp1": tp1,
        "plan_tp2": tp2,
        "plan_risk": risk,
        "plan_tp1_r": tp1_r,
        "plan_tp2_r": tp2_r,
        "plan_sl_atr_buffer": sl_atr_buffer,
        "plan_sl_valid": sl_valid,
        "plan_disp_index": disp.get("index"),
        "plan_retest_index": rt.get("index"),
        "plan_atr_abs": atr_f,
    }


def walk_plan_outcome(
    candles: list[Candle],
    start_index: int,
    direction: str,
    sl: float,
    tp1: float,
    tp2: float,
) -> dict[str, Any]:
    """Forward-walk OHLC after the RETEST bar. SL before TP on the same bar."""
    for c in candles[start_index + 1 :]:
        if direction == "LONG":
            hit_sl = c.low <= sl
            hit_tp1 = c.high >= tp1
            hit_tp2 = c.high >= tp2
        else:
            hit_sl = c.high >= sl
            hit_tp1 = c.low <= tp1
            hit_tp2 = c.low <= tp2
        if hit_sl:
            return {
                "plan_outcome": "SL_HIT",
                "plan_exit_bar": c.index,
                "plan_exit_ts": c.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                "plan_exit_price": sl,
            }
        if hit_tp2:
            return {
                "plan_outcome": "TP2_HIT",
                "plan_exit_bar": c.index,
                "plan_exit_ts": c.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                "plan_exit_price": tp2,
            }
        if hit_tp1:
            return {
                "plan_outcome": "TP1_HIT",
                "plan_exit_bar": c.index,
                "plan_exit_ts": c.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                "plan_exit_price": tp1,
            }
    last = candles[-1]
    return {
        "plan_outcome": "OPEN_AT_END",
        "plan_exit_bar": last.index,
        "plan_exit_ts": last.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
        "plan_exit_price": last.close,
    }


# Research grids: x = tp1_atr_multiplier, y = tp2_atr_multiplier, z = sl_atr_buffer.
# Geometry-only — CRT states are not replayed per cell (SL/TP do not change the machine).
DEFAULT_X_TP1 = (0.5, 1.0, 1.5, 2.0, 3.0)
DEFAULT_Y_TP2 = (1.0, 2.0, 3.0, 4.0)
DEFAULT_Z_SL_BUF = (0.1, 0.2, 0.3, 0.5)

# Feature columns stamped on each setup at its own bar (config is constant; values are not).
SETUP_FEATURE_WATCH = (
    "body_ratio", "atr", "ema_spread", "momentum_score", "trend_strength",
    "rsi_14", "session", "hour_of_day", "retest_depth", "disp_strength",
    "displacement_retrace", "displacement_atr_ratio", "candles_since_retest",
    "volatility_ratio", "volume_ratio",
)


def _parse_float_grid(raw: str, default: tuple[float, ...]) -> tuple[float, ...]:
    if raw is None or str(raw).strip() == "":
        return default
    vals = tuple(float(p.strip()) for p in str(raw).split(",") if p.strip())
    if not vals:
        raise ValueError("empty research grid")
    return vals


def _features_at_bar(feat_df: pd.DataFrame, bar_index: int) -> dict[str, Any]:
    row = feat_df.iloc[bar_index]
    out: dict[str, Any] = {
        "warmup_complete": bool(row["warmup_complete"]),
    }
    for col in CANONICAL_FEATURES:
        if col not in feat_df.columns:
            out[col] = None
            continue
        val = row[col]
        out[col] = None if pd.isna(val) else (float(val) if col != "session" else val)
    # session may be numeric ordinal or label depending on pipeline emission
    if "session" in feat_df.columns and not pd.isna(row["session"]):
        try:
            out["session"] = float(row["session"])
        except (TypeError, ValueError):
            out["session"] = row["session"]
    return out


def run_xy_research(
    candles: list[Candle],
    crt_records: list[dict[str, Any]],
    feat_df: pd.DataFrame,
    fp_cfg: dict[str, Any],
    *,
    x_grid: tuple[float, ...],
    y_grid: tuple[float, ...],
    z_grid: tuple[float, ...],
) -> dict[str, Any]:
    """Sweep x=TP1, y=TP2, z=sl_atr_buffer over RETEST_CONFIRMED bars.

    One CRT trajectory. Feature-pipeline config is the run's live snapshot (does not
    vary by bar). Per-setup feature *values* are taken from that bar's timestamp.
    """
    setups = [r for r in crt_records if r.get("action") == "RETEST_CONFIRMED"]
    setup_meta: list[dict[str, Any]] = []
    for rec in setups:
        feats = _features_at_bar(feat_df, rec["bar_index"])
        watch = {k: feats.get(k) for k in SETUP_FEATURE_WATCH}
        setup_meta.append(
            {
                "bar_index": rec["bar_index"],
                "timestamp": rec["timestamp"],
                "direction": rec.get("direction"),
                "crt_action": rec.get("action"),
                "crt_reject_reason": rec.get("reject_reason"),
                "state_after": rec.get("state_after"),
                "crt_atr_abs": rec.get("crt_atr_abs"),
                "feature_values_at_bar": watch,
                "canonical_features_at_bar": feats,
                "feature_config_at_bar": {
                    "note": (
                        "Feature-pipeline config is production-current for the whole book; "
                        "there is no per-bar historical config version. Values below are "
                        "the live FeaturePipeline._fp_cfg that emitted this bar."
                    ),
                    "identity_selectors": {
                        "normalization_basis": fp_cfg["normalization_basis"],
                        "session_timestamp_basis": fp_cfg["session_timestamp_basis"],
                        "swing_window": fp_cfg["swing_window"],
                    },
                    "config": fp_cfg,
                },
            }
        )

    rows: list[dict[str, Any]] = []
    for rec, meta in zip(setups, setup_meta):
        for x in x_grid:
            for y in y_grid:
                for z in z_grid:
                    plan = planned_entry_sl_tp(
                        rec, tp1_r=float(x), tp2_r=float(y), sl_atr_buffer=float(z)
                    )
                    cell: dict[str, Any] = {
                        "timestamp": rec["timestamp"],
                        "bar_index": rec["bar_index"],
                        "direction": rec.get("direction"),
                        "x_tp1": float(x),
                        "y_tp2": float(y),
                        "z_sl_atr_buffer": float(z),
                        "xy_order_ok": float(x) < float(y),
                    }
                    if plan is None:
                        cell["plan_outcome"] = "NO_PLAN"
                        cell["r_if_all_in"] = None
                        rows.append(cell)
                        continue
                    cell.update(
                        {
                            "plan_entry": plan["plan_entry"],
                            "plan_sl": plan["plan_sl"],
                            "plan_tp1": plan["plan_tp1"],
                            "plan_tp2": plan["plan_tp2"],
                            "plan_risk": plan["plan_risk"],
                            "plan_sl_valid": plan["plan_sl_valid"],
                            "plan_atr_abs": plan["plan_atr_abs"],
                        }
                    )
                    if not plan["plan_sl_valid"]:
                        cell["plan_outcome"] = "INVERTED_SL"
                        cell["r_if_all_in"] = 0.0
                        rows.append(cell)
                        continue
                    walked = walk_plan_outcome(
                        candles,
                        rec["bar_index"],
                        plan["plan_direction"],
                        plan["plan_sl"],
                        plan["plan_tp1"],
                        plan["plan_tp2"],
                    )
                    cell.update(walked)
                    outcome = walked["plan_outcome"]
                    if outcome == "SL_HIT":
                        cell["r_if_all_in"] = -1.0
                    elif outcome == "TP1_HIT":
                        cell["r_if_all_in"] = float(x)
                    elif outcome == "TP2_HIT":
                        cell["r_if_all_in"] = float(y)
                    else:
                        cell["r_if_all_in"] = None
                    rows.append(cell)

    summary_cells: list[dict[str, Any]] = []
    for x in x_grid:
        for y in y_grid:
            for z in z_grid:
                subset = [
                    r for r in rows
                    if r["x_tp1"] == float(x)
                    and r["y_tp2"] == float(y)
                    and r["z_sl_atr_buffer"] == float(z)
                ]
                outcomes: dict[str, int] = {}
                rs = [r["r_if_all_in"] for r in subset if r.get("r_if_all_in") is not None]
                for r in subset:
                    outcomes[r["plan_outcome"]] = outcomes.get(r["plan_outcome"], 0) + 1
                summary_cells.append(
                    {
                        "x_tp1": float(x),
                        "y_tp2": float(y),
                        "z_sl_atr_buffer": float(z),
                        "xy_order_ok": float(x) < float(y),
                        "n_setups": len(subset),
                        "outcomes": outcomes,
                        "sum_r_if_all_in": (sum(rs) if rs else None),
                        "mean_r_if_all_in": (sum(rs) / len(rs) if rs else None),
                    }
                )
    summary_cells.sort(
        key=lambda c: (
            c["mean_r_if_all_in"] is None,
            -(c["mean_r_if_all_in"] or 0.0),
            c["x_tp1"],
            c["y_tp2"],
            c["z_sl_atr_buffer"],
        )
    )
    return {
        "variables": {
            "x": "tp1_atr_multiplier (R-multiple of |entry-SL|)",
            "y": "tp2_atr_multiplier (R-multiple of |entry-SL|)",
            "z": "sl_atr_buffer (ATR units beyond displacement extreme)",
        },
        "grids": {"x_tp1": list(x_grid), "y_tp2": list(y_grid), "z_sl_atr_buffer": list(z_grid)},
        "caveat": (
            "r_if_all_in is all-in at first touch (SL=-1, TP1=+x, TP2=+y). "
            "It is NOT the live 0.5/0.5 partial blend and NOT expectancy. "
            "N is the RETEST_CONFIRMED count on this book — not a promotion sample. "
            "F-025: exit/cost is not an expectancy lever."
        ),
        "setups": setup_meta,
        "cells": summary_cells,
        "rows": rows,
    }


# ─────────────────────────────────────────────────────────────────────────────
# 4. GATE OVERLAY  (EngineRunner -> fusion -> decision; NOT fed back)
# ─────────────────────────────────────────────────────────────────────────────
def _merge_engine_config(prod: dict) -> dict:
    """Mirror live_engine_hook._load_engine_config flatten (read-only process copy).

    Adapted from reports/xauusd_retest_pathb_counterfactual_probe.py:179-199, WITHOUT its
    force_session_pass override — production session policy applies here.
    """
    engine_cfg = dict(prod["engine_runner"])
    merged: dict[str, Any] = {}
    merged.update(engine_cfg)
    merged.update(dict(prod["decision_engine"]))
    merged["fusion_engine"] = dict(prod["fusion_engine"])
    merged["ultron_risk_gate"] = dict(prod["ultron_risk_gate"])
    merged["execution_planner"] = dict(prod["execution_planner"])
    merged["crt_engine"] = dict(prod["crt_engine"])
    if "rr_fusion" not in merged and "rr_fusion" in engine_cfg:
        merged["rr_fusion"] = engine_cfg["rr_fusion"]
    return merged


_EXECUTE_TOKENS = {"execute", "accept", "go", "approved", "approve"}


def run_gate_overlay(
    feat_df: pd.DataFrame,
    crt_records: list[dict[str, Any]],
    mode: str,
    instrument: str,
) -> list[dict[str, Any]]:
    """Read-only fusion overlay. Returns one dict per bar (empty dict = not evaluated)."""
    blanks: list[dict[str, Any]] = [{} for _ in crt_records]
    if mode == "none":
        return blanks

    prod = get_prod_metadata()
    runner = EngineRunner(_merge_engine_config(prod))
    feat_cols = [c for c in CANONICAL_FEATURES if c in feat_df.columns]

    evaluated = 0
    for i, rec in enumerate(crt_records):
        if not rec["warmup_complete_hint"]:
            continue
        d_int = {"LONG": 1, "SHORT": -1}.get(rec.get("direction") or "", 0)
        if mode == "events" and d_int == 0 and rec.get("action") in (None, "NONE"):
            continue

        row = feat_df.iloc[i]
        input_data: dict[str, Any] = {c: float(row[c]) for c in feat_cols}
        for k in ("open", "high", "low", "close", "volume"):
            input_data[k] = float(row[k])
        input_data["_data_integrity"] = "real"
        input_data["instrument"] = instrument
        input_data["timestamp"] = rec["timestamp"]
        input_data["bar_id"] = str(rec["bar_index"])
        input_data["direction"] = d_int
        input_data["signal_dir"] = d_int
        input_data["trade_direction"] = d_int
        # FeaturePipeline's `atr` slot is close-relative; engines want price units (FM-074).
        atr_rel = float(input_data.get("atr") or 0.0)
        close = float(input_data["close"])
        if 0 < atr_rel < 1.0 and close > 0:
            input_data["atr_relative"] = atr_rel
            input_data["atr"] = atr_rel * close

        context = {
            "instrument": instrument,
            "timeframe": TIMEFRAME,
            "symbol": instrument,
            "strategy_consensus_direction": d_int,
        }
        try:
            er = runner.run(input_data, context=context)
        except Exception as exc:  # observe-only: an overlay error must not kill the trace
            blanks[i] = {"gate_error": str(exc)[:200]}
            continue

        decision = str(er.get("decision") or er.get("status") or "")
        blanks[i] = {
            "fusion_score": er.get("final_score", er.get("score")),
            "decision": decision,
            "confidence": er.get("confidence"),
            "veto_reason": er.get("reason"),
            "reject_stage": er.get("reject_stage"),
            "would_veto": decision.lower() not in _EXECUTE_TOKENS,
        }
        evaluated += 1

    logger.info("gate overlay evaluated %d/%d bars (mode=%s)", evaluated, len(crt_records), mode)
    return blanks


# ─────────────────────────────────────────────────────────────────────────────
# 5. OUTPUT
# ─────────────────────────────────────────────────────────────────────────────
GATE_COLS = [
    "fusion_score", "decision", "confidence", "veto_reason",
    "reject_stage", "would_veto", "gate_error",
]
STATE_COLS = [
    "warmup_complete", "htf_seeded", "state_before", "state_after", "action", "direction",
    "risk_score", "crt_atr_abs", "ema_fast_val", "ema_slow_val", "range_session",
    "event", "reject_reason",
    "trade_entry", "trade_sl", "trade_tp1", "trade_tp2", "trade_risk_pct", "trade_id",
    "plan_entry", "plan_sl", "plan_tp1", "plan_tp2", "plan_risk",
    "plan_tp1_r", "plan_tp2_r", "plan_sl_atr_buffer", "plan_sl_valid",
    "plan_outcome", "plan_exit_bar", "plan_exit_ts", "plan_exit_price",
]


def write_outputs(
    instrument: str,
    out_dir: Path,
    feat_df: pd.DataFrame,
    crt_records: list[dict[str, Any]],
    gate_rows: list[dict[str, Any]],
    manifest: dict[str, Any],
) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)

    # ── trace.csv — the flat joined table ──────────────────────────────────
    flat = pd.DataFrame(
        [{k: v for k, v in r.items() if not k.startswith("_")} for r in crt_records]
    )
    gate_df = pd.DataFrame(gate_rows)
    for col in GATE_COLS:
        if col not in gate_df.columns:
            gate_df[col] = np.nan

    ordered = ["timestamp", "bar_index", "open", "high", "low", "close", "volume"]
    csv = pd.concat(
        [
            feat_df[["timestamp", "open", "high", "low", "close", "volume"]].reset_index(drop=True),
            feat_df[list(CANONICAL_FEATURES)].reset_index(drop=True),
            flat.drop(columns=["timestamp"], errors="ignore").reset_index(drop=True),
            gate_df[GATE_COLS].reset_index(drop=True),
        ],
        axis=1,
    )
    csv["timestamp"] = [r["timestamp"] for r in crt_records]
    cols = (
        ordered
        + list(CANONICAL_FEATURES)
        + [c for c in STATE_COLS if c in csv.columns]
        + GATE_COLS
    )
    csv = csv[[c for c in cols if c in csv.columns]]
    csv.to_csv(out_dir / "trace.csv", index=False, na_rep="")

    # ── trace.jsonl — same bars, with the nested crt_inputs block ──────────
    with open(out_dir / "trace.jsonl", "w", encoding="utf-8") as fh:
        for i, rec in enumerate(crt_records):
            row = feat_df.iloc[i]
            payload = {
                "symbol": instrument,
                "timestamp": rec["timestamp"],
                "bar_index": rec["bar_index"],
                "raw_ohlcv": {
                    k: float(row[k]) for k in ("open", "high", "low", "close", "volume")
                },
                "features": {
                    c: (None if pd.isna(row[c]) else float(row[c]))
                    for c in CANONICAL_FEATURES if c in feat_df.columns
                },
                "warmup_complete": bool(row["warmup_complete"]),
                "crt": {
                    k: v for k, v in rec.items()
                    if not k.startswith("_") and k not in ("timestamp", "bar_index")
                },
                "crt_inputs": _plain(rec["_crt_inputs"]),
                "gate_overlay": _plain(gate_rows[i]) or None,
            }
            fh.write(json.dumps(payload, default=str) + "\n")

    with open(out_dir / "manifest.json", "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, indent=2, default=str)

    knobs = manifest.get("configurable") or []
    with open(out_dir / "knobs.json", "w", encoding="utf-8") as fh:
        json.dump(knobs, fh, indent=2, default=str)

    feat_block = manifest.get("features") or {}
    with open(out_dir / "feature_config.json", "w", encoding="utf-8") as fh:
        json.dump(
            {
                "identity_selectors": feat_block.get("identity_selectors"),
                "derived_warmup_rows": feat_block.get("derived_warmup_rows"),
                "config_source": feat_block.get("config_source"),
                "config_untouched_by_htf": feat_block.get("config_untouched_by_htf"),
                "config": feat_block.get("config"),
            },
            fh,
            indent=2,
            default=str,
        )

    setups = [r for r in crt_records if r.get("action") in ("RETEST_CONFIRMED", "TRADE_OPENED")]
    if setups:
        keep = [
            "bar_index", "timestamp", "action", "direction",
            "plan_entry", "plan_sl", "plan_tp1", "plan_tp2", "plan_risk",
            "plan_tp1_r", "plan_tp2_r", "plan_sl_atr_buffer", "plan_sl_valid",
            "plan_outcome", "plan_exit_bar", "plan_exit_ts", "plan_exit_price",
            "plan_disp_index", "plan_retest_index", "plan_atr_abs",
            "trade_entry", "trade_sl", "trade_tp1", "trade_tp2", "trade_id",
            "reject_reason",
        ]
        rows = [{k: s.get(k) for k in keep} for s in setups]
        pd.DataFrame(rows).to_csv(out_dir / "setups.csv", index=False, na_rep="")

    xy = manifest.get("xy_research")
    if xy:
        slim_rows = [
            {k: r.get(k) for k in (
                "timestamp", "bar_index", "direction",
                "x_tp1", "y_tp2", "z_sl_atr_buffer", "xy_order_ok",
                "plan_entry", "plan_sl", "plan_tp1", "plan_tp2", "plan_risk",
                "plan_sl_valid", "plan_outcome", "plan_exit_ts", "r_if_all_in",
            )}
            for r in xy.get("rows", [])
        ]
        pd.DataFrame(slim_rows).to_csv(out_dir / "research_xy.csv", index=False, na_rep="")
        setups_feat = []
        for s in xy.get("setups", []):
            setups_feat.append(
                {
                    "bar_index": s["bar_index"],
                    "timestamp": s["timestamp"],
                    "direction": s["direction"],
                    "crt_atr_abs": s["crt_atr_abs"],
                    "crt_reject_reason": s.get("crt_reject_reason"),
                    "feature_config_identity": s["feature_config_at_bar"]["identity_selectors"],
                    **{f"feat_{k}": v for k, v in s["feature_values_at_bar"].items()},
                }
            )
        pd.DataFrame(setups_feat).to_csv(
            out_dir / "research_xy_setup_features.csv", index=False, na_rep=""
        )
        with open(out_dir / "research_xy_summary.json", "w", encoding="utf-8") as fh:
            payload = {k: v for k, v in xy.items() if k != "rows"}
            # Drop full 39-dim from summary JSON; CSV + trace.jsonl keep values.
            slim_setups = []
            for s in payload.get("setups", []):
                slim = dict(s)
                slim.pop("canonical_features_at_bar", None)
                slim_setups.append(slim)
            payload["setups"] = slim_setups
            json.dump(payload, fh, indent=2, default=str)


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group()
    src.add_argument("--xlsx", type=Path, default=None, help=f"MT5 workbook (default: {DEFAULT_XLSX.relative_to(ROOT)})")
    src.add_argument("--csv", type=Path, default=None, help="CSV alternative with the same OHLCV columns")
    ap.add_argument("--instrument", default=DEFAULT_INSTRUMENT,
                    help="instrument symbol; also picks the default corpus data/mt5/<INST>_M15.csv")
    ap.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    ap.add_argument(
        "--overlay", choices=("all", "events", "none"), default="all",
        help="Bars to run the fusion overlay on: all bars, only CRT-active bars, or skip.",
    )
    ap.add_argument(
        "--htf-bars", type=int, default=None,
        help="M15 bars per HTF window. Default: production backtest.htf_candles_per_range. "
             "This is the only knob this flag touches — it does not change shadow TTL.",
    )
    ap.add_argument(
        "--tp1-r", type=float, default=None,
        help="Override TP1 as R-multiple of |entry-SL|. Production default 1.0 (intent-specific).",
    )
    ap.add_argument(
        "--tp2-r", type=float, default=None,
        help="Override TP2 as R-multiple of |entry-SL|. Production default 2.0.",
    )
    ap.add_argument(
        "--sl-atr-buffer", type=float, default=None,
        help="Override SL ATR buffer. Production default 0.2.",
    )
    ap.add_argument(
        "--shadow-ttl", type=int, default=None,
        help="Override pending_displacement_ttl_candles. Default: production value (independent of --htf-bars).",
    )
    ap.add_argument(
        "--ignore-off-session",
        action="store_true",
        help=(
            "This-run only: add OFF_SESSION to CRTConfig.allowed_sessions so the "
            "session gate does not FILTER_REJECT. Does not rewrite production JSON "
            "or change session_windows / feature session_timestamp_basis."
        ),
    )
    ap.add_argument(
        "--research-xy",
        action="store_true",
        help=(
            "After the CRT replay, sweep x=tp1_atr_multiplier, y=tp2_atr_multiplier, "
            "z=sl_atr_buffer on RETEST_CONFIRMED bars. Does not change CRT states or "
            "write production JSON."
        ),
    )
    ap.add_argument("--tp1-grid", default=None, help="Comma list for x. Default: 0.5,1.0,1.5,2.0,3.0")
    ap.add_argument("--tp2-grid", default=None, help="Comma list for y. Default: 1.0,2.0,3.0,4.0")
    ap.add_argument(
        "--sl-buf-grid", default=None, help="Comma list for z sl_atr_buffer. Default: 0.1,0.2,0.3,0.5"
    )
    args = ap.parse_args(argv or sys.argv[1:])

    _assert_stage_list_matches_run_source()
    instrument = args.instrument.upper()

    # Explicit --csv/--xlsx wins; otherwise the instrument picks its canonical MT5 corpus.
    # DEFAULT_XLSX is only the fallback for the instrument it belongs to.
    if args.csv or args.xlsx:
        source = (args.csv or args.xlsx).resolve()
    elif instrument == DEFAULT_INSTRUMENT and DEFAULT_XLSX.exists():
        source = DEFAULT_XLSX
    else:
        source = corpus_for(instrument)
    source = Path(source).resolve()
    if not source.exists():
        logger.error("source not found for %s: %s", instrument, source)
        return 1

    if source.suffix.lower() in (".xlsx", ".xlsm"):
        candles, source_info = load_from_xlsx(source)
    else:
        candles, source_info = load_from_csv(source)
    logger.info("loaded %d candles from %s", len(candles), source)

    feat_df, fp_cfg_used = build_feature_frame(candles)
    logger.info(
        "features: %d/%d bars warmup-complete", int(feat_df["warmup_complete"].sum()), len(feat_df)
    )

    prod_cfg = load_prod_config_from_registry(get_active_version(), instrument)
    prod_htf = int(get_prod_section("backtest")["htf_candles_per_range"])
    htf_bars = prod_htf if args.htf_bars is None else int(args.htf_bars)
    overrides: dict[str, Any] = {}
    tp1_r = float(prod_cfg.tp1_atr_multiplier) if args.tp1_r is None else float(args.tp1_r)
    tp2_r = float(prod_cfg.tp2_atr_multiplier) if args.tp2_r is None else float(args.tp2_r)
    sl_buf = float(prod_cfg.sl_atr_buffer) if args.sl_atr_buffer is None else float(args.sl_atr_buffer)
    if args.tp1_r is not None:
        overrides.update(
            {
                "tp1_atr_multiplier": tp1_r,
                "tp1_atr_multiplier_breakout": tp1_r,
                "tp1_atr_multiplier_pullback": tp1_r,
                "tp1_atr_multiplier_liq_sweep": tp1_r,
                "tp1_atr_multiplier_reversal": tp1_r,
            }
        )
    if args.tp2_r is not None:
        overrides["tp2_atr_multiplier"] = tp2_r
    if args.sl_atr_buffer is not None:
        overrides["sl_atr_buffer"] = sl_buf
    if args.shadow_ttl is not None:
        shadow_ttl = int(args.shadow_ttl)
        overrides["pending_displacement_ttl_candles"] = shadow_ttl
    else:
        shadow_ttl = int(prod_cfg.pending_displacement_ttl_candles)
    if args.ignore_off_session:
        overrides["allowed_sessions"] = tuple(
            dict.fromkeys((*prod_cfg.allowed_sessions, "OFF_SESSION"))
        )
    crt_cfg = replace(prod_cfg, **overrides) if overrides else prod_cfg

    crt_records, engine = replay_crt(
        candles, instrument, htf_bars=htf_bars, crt_cfg=crt_cfg
    )
    for i, rec in enumerate(crt_records):
        rec["warmup_complete"] = bool(feat_df.iloc[i]["warmup_complete"])
        rec["warmup_complete_hint"] = rec["warmup_complete"]
        if rec.get("action") in ("RETEST_CONFIRMED", "TRADE_OPENED"):
            plan = planned_entry_sl_tp(
                rec, tp1_r=tp1_r, tp2_r=tp2_r, sl_atr_buffer=sl_buf
            )
            if plan:
                rec.update(plan)
                rec.update(
                    walk_plan_outcome(
                        candles,
                        rec["bar_index"],
                        plan["plan_direction"],
                        plan["plan_sl"],
                        plan["plan_tp1"],
                        plan["plan_tp2"],
                    )
                )

    gate_rows = run_gate_overlay(feat_df, crt_records, args.overlay, instrument)
    for rec in crt_records:
        rec.pop("warmup_complete_hint", None)

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    out_dir = args.output_dir / stamp

    state_counts: dict[str, int] = {}
    action_counts: dict[str, int] = {}
    for rec in crt_records:
        if rec.get("state_after"):
            state_counts[rec["state_after"]] = state_counts.get(rec["state_after"], 0) + 1
        if rec.get("action"):
            action_counts[rec["action"]] = action_counts.get(rec["action"], 0) + 1

    xy_block: dict[str, Any] | None = None
    if args.research_xy:
        xy_block = run_xy_research(
            candles,
            crt_records,
            feat_df,
            fp_cfg_used,
            x_grid=_parse_float_grid(args.tp1_grid, DEFAULT_X_TP1),
            y_grid=_parse_float_grid(args.tp2_grid, DEFAULT_Y_TP2),
            z_grid=_parse_float_grid(args.sl_buf_grid, DEFAULT_Z_SL_BUF),
        )

    manifest = {
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "script": "scripts/analysis/xauusd_excel_feature_state_trace.py",
        "read_only": True,
        "source": {
            "path": str(source.relative_to(ROOT)) if source.is_relative_to(ROOT) else str(source),
            "sha256": _sha256(source),
            "provenance": _plain(source_info) or None,
        },
        "corpus": {
            "instrument": instrument,
            "timeframe": TIMEFRAME,
            "bars": len(candles),
            "first_timestamp": crt_records[0]["timestamp"] if crt_records else None,
            "last_timestamp": crt_records[-1]["timestamp"] if crt_records else None,
            "timestamp_basis": "broker-server time as exported by MT5 (see F-066)",
        },
        "config": {
            "active_version": get_active_version(),
            "schema_version": SCHEMA_VERSION,
            "canonical_features": len(CANONICAL_FEATURES),
            "backtest_engine_gate_env": os.environ.get("BACKTEST_ENGINE_GATE"),
        },
        "configurable": [
            {
                "knob": "htf_bars",
                "used": htf_bars,
                "production_default": prod_htf,
                "source": "cli" if args.htf_bars is not None else "production",
                "where": "backtest.htf_candles_per_range → HTFBuilder(n). CRTConfig does not own this.",
                "meaning": (
                    "M15 bars per completed HTF window. Flip of current_htf_id is the only "
                    "direct effect. Does not change retrace, TTL, sessions, or TP."
                ),
            },
            {
                "knob": "tp1_r",
                "used": tp1_r,
                "production_default": float(prod_cfg.tp1_atr_multiplier),
                "source": "cli" if args.tp1_r is not None else "production",
                "where": "CRTConfig.tp1_atr_multiplier (+ intent-specific keys when --tp1-r set)",
                "meaning": "TP1 distance as R-multiple of |entry-SL|.",
            },
            {
                "knob": "tp2_r",
                "used": tp2_r,
                "production_default": float(prod_cfg.tp2_atr_multiplier),
                "source": "cli" if args.tp2_r is not None else "production",
                "where": "CRTConfig.tp2_atr_multiplier",
                "meaning": "TP2 / runner distance as R-multiple of |entry-SL|.",
            },
            {
                "knob": "sl_atr_buffer",
                "used": sl_buf,
                "production_default": float(prod_cfg.sl_atr_buffer),
                "source": "cli" if args.sl_atr_buffer is not None else "production",
                "where": "CRTConfig.sl_atr_buffer",
                "meaning": "SL offset beyond displacement extreme, in engine ATR units.",
            },
            {
                "knob": "shadow_ttl",
                "used": shadow_ttl,
                "production_default": int(prod_cfg.pending_displacement_ttl_candles),
                "source": "cli" if args.shadow_ttl is not None else "production",
                "where": "CRTConfig.pending_displacement_ttl_candles",
                "meaning": "Bars a DISPLACEMENT memory survives after an HTF reset. Independent of htf_candles_per_range.",
            },
            {
                "knob": "allowed_sessions",
                "used": list(crt_cfg.allowed_sessions),
                "production_default": list(prod_cfg.allowed_sessions),
                "source": "cli" if args.ignore_off_session else "production",
                "where": "CRTConfig.allowed_sessions",
                "meaning": (
                    "Sessions that may open a trade. --ignore-off-session adds OFF_SESSION "
                    "only; session_windows and feature session_timestamp_basis stay production."
                ),
            },
            {
                "knob": "session_windows",
                "used": {k: [str(a), str(b)] for k, (a, b) in prod_cfg.session_windows.items()},
                "source": "production",
                "where": "CRTConfig.session_windows",
                "meaning": "Clock windows for session names. Broker-server time on this corpus.",
            },
        ],
        "features": {
            "stages_run": PIPELINE_STAGES,
            "finalize_called": False,
            "warmup_policy": "NaN left empty; never ffill/bfill/zero-filled",
            "warmup_complete_bars": int(feat_df["warmup_complete"].sum()),
            "derived_warmup_rows": required_warmup_rows(fp_cfg_used),
            "config_source": "FeaturePipeline._fp_cfg via get_prod_section('feature_pipeline')",
            "config_untouched_by_htf": True,
            "identity_selectors": {
                "normalization_basis": fp_cfg_used["normalization_basis"],
                "session_timestamp_basis": fp_cfg_used["session_timestamp_basis"],
                "swing_window": fp_cfg_used["swing_window"],
            },
            "config": fp_cfg_used,
        },
        "states": {"state_after_counts": state_counts, "action_counts": action_counts},
        "xy_research": xy_block,
        "gate_overlay": {
            "mode": args.overlay,
            "evaluated_bars": sum(1 for g in gate_rows if g),
            "would_veto_bars": sum(1 for g in gate_rows if g.get("would_veto")),
            "caveat": (
                "Read-only overlay: the veto is NOT fed back into the CRT state machine, so both "
                "passes share one trajectory. A real gate-ON replay diverges (F-055: a reject "
                "resets the state machine, removed != added). Answers 'which committed entries "
                "would fusion veto', not 'what would the ledger look like'. F-070 measured 0/30 "
                "vetoes on the active config."
            ),
        },
    }

    write_outputs(instrument, out_dir, feat_df, crt_records, gate_rows, manifest)
    logger.info("wrote -> %s", _rel(out_dir))
    print(f"\ntrace.csv    : {_rel(out_dir / 'trace.csv')}")
    print(f"trace.jsonl  : {_rel(out_dir / 'trace.jsonl')}")
    print(f"setups.csv   : {_rel(out_dir / 'setups.csv')}")
    print(f"knobs.json   : {_rel(out_dir / 'knobs.json')}")
    print(f"feature_config.json: {_rel(out_dir / 'feature_config.json')}")
    if xy_block is not None:
        print(f"research_xy.csv: {_rel(out_dir / 'research_xy.csv')}")
        print(f"research_xy_setup_features.csv: {_rel(out_dir / 'research_xy_setup_features.csv')}")
        print(f"research_xy_summary.json: {_rel(out_dir / 'research_xy_summary.json')}")
    print(f"manifest.json: {_rel(out_dir / 'manifest.json')}")
    print(
        f"\nbars={len(candles)}  htf_bars={htf_bars}  tp1_r={tp1_r}  tp2_r={tp2_r}  "
        f"ignore_off_session={args.ignore_off_session}  "
        f"allowed_sessions={list(crt_cfg.allowed_sessions)}"
    )
    print(
        "feature_identity "
        f"normalization_basis={fp_cfg_used['normalization_basis']} "
        f"session_timestamp_basis={fp_cfg_used['session_timestamp_basis']} "
        f"swing_window={fp_cfg_used['swing_window']} "
        f"warmup_rows={required_warmup_rows(fp_cfg_used)}"
    )
    print(f"states={state_counts}\nactions={action_counts}")
    if xy_block is not None:
        print(
            f"xy_research setups={len(xy_block['setups'])} "
            f"cells={len(xy_block['cells'])} "
            f"x={xy_block['grids']['x_tp1']} y={xy_block['grids']['y_tp2']} "
            f"z={xy_block['grids']['z_sl_atr_buffer']}"
        )
        top = [c for c in xy_block["cells"] if c.get("xy_order_ok")][:5]
        print("top xy_order_ok cells by mean_r_if_all_in:")
        for c in top:
            print(
                f"  x={c['x_tp1']} y={c['y_tp2']} z={c['z_sl_atr_buffer']} "
                f"mean_r={c['mean_r_if_all_in']} outcomes={c['outcomes']}"
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
