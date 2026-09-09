"""Inventory mt5_analytics / oss_lab / tools → .grok/gcmc_v2_inventory.xlsx (GCMC v2)."""
from __future__ import annotations

import ast
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
TREES = ("mt5_analytics", "oss_lab", "tools")
OUT = ROOT / ".grok" / "gcmc_v2_inventory.xlsx"
SKIP = {"__pycache__", ".git", "venv", ".venv"}


def summary(path: Path) -> str:
    try:
        src = path.read_text(encoding="utf-8", errors="replace")
    except OSError as e:
        return f"(unreadable: {e})"
    try:
        tree = ast.parse(src)
        doc = ast.get_docstring(tree) or ""
    except SyntaxError as e:
        doc = ""
        m = re.search(r'"""(.*?)"""', src, re.S)
        if m:
            doc = m.group(1)
        else:
            return f"(syntax: {e})"
    lines = [ln.strip() for ln in doc.splitlines() if ln.strip()]
    if not lines:
        return f"Module {path.stem}."
    text = re.sub(r"\s+", " ", " ".join(lines[:4]))
    return text[:400]


def main() -> None:
    rows: list[tuple[str, str, str]] = []
    for tree in TREES:
        root = ROOT / tree
        if not root.exists():
            continue
        for p in sorted(root.rglob("*.py")):
            if any(x in p.parts for x in SKIP):
                continue
            rel = p.relative_to(ROOT).as_posix()
            rows.append((tree, rel, summary(p)))

    wb = Workbook()
    ws = wb.active
    ws.title = "gcmc_v2"
    headers = ["Tree", "File Name", "Summary of functionality"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 90
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    counts = ws_counts = wb.create_sheet("Counts")
    from collections import Counter

    c = Counter(r[0] for r in rows)
    counts.append(["Metric", "Value"])
    counts.append(["GCMC v2 listed", len(rows)])
    for t in TREES:
        counts.append([f"{t} listed", c.get(t, 0)])
    wb.save(OUT)
    print(f"wrote {len(rows)} rows -> {OUT}")
    for t in TREES:
        print(f"  {t}: {c.get(t, 0)}")


if __name__ == "__main__":
    main()
