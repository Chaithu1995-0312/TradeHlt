"""One-shot exporter: scripts/**/*.py -> Excel (name, summary, project imports)."""
from __future__ import annotations

import ast
import re
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

ROOT = Path(__file__).resolve().parent
SCRIPTS = ROOT / "scripts"
OUT = ROOT / "scripts_business_functionality.xlsx"

PROJECT_TOPS = {
    "src",
    "scripts",
    "configs",
    "tests",
    "models",
    "multi_llm",
    "core",
    "engines",
    "config_layer",
    "features",
    "governance",
    "runtime",
    "strategies",
    "training",
    "agent",
    "bitnet",
    "control_plane",
    "inout",
    "utils",
    "research",
    "interpreters",
    "portfolio",
}


def first_meaningful_summary(text: str, max_len: int = 500) -> str:
    if not text:
        return ""
    t = text.strip()
    lines = [ln.rstrip() for ln in t.splitlines()]
    cleaned: list[str] = []
    for ln in lines:
        s = ln.strip()
        if not s:
            if cleaned:
                break
            continue
        if re.fullmatch(r"[=#\-*_~]{3,}", s):
            continue
        if re.fullmatch(r"[\w./\-]+\.py\b.*", s) and len(s) < 80:
            continue
        cleaned.append(s)
    para = re.sub(r"\s+", " ", " ".join(cleaned)).strip()
    parts = re.split(r"(?<=[.!?])\s+", para)
    summary = " ".join(parts[:2]).strip() if parts else para
    if len(summary) > max_len:
        summary = summary[: max_len - 1].rsplit(" ", 1)[0] + "…"
    return summary


def summarize_from_filename(rel: str) -> str:
    name = Path(rel).stem
    name = re.sub(r"^_+", "", name)
    words = re.sub(r"[_\-]+", " ", name).strip()
    return f"Script utility: {words}."


def try_argparse_description(tree: ast.AST) -> str:
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call):
            continue
        func = node.func
        name = ""
        if isinstance(func, ast.Name):
            name = func.id
        elif isinstance(func, ast.Attribute):
            name = func.attr
        if name != "ArgumentParser":
            continue
        for kw in node.keywords:
            if (
                kw.arg == "description"
                and isinstance(kw.value, ast.Constant)
                and isinstance(kw.value.value, str)
            ):
                return first_meaningful_summary(kw.value.value, 400)
    return ""


def resolve_module_to_files(mod: str) -> list[str]:
    if not mod:
        return []
    top = mod.split(".")[0]
    parts = mod.split(".")
    candidates: list[Path] = []
    candidates.append(ROOT.joinpath(*parts))
    candidates.append(ROOT.joinpath(*parts).with_suffix(".py"))
    for base in (ROOT / "src", ROOT, ROOT / "scripts"):
        candidates.append(base.joinpath(*parts))
        candidates.append(base.joinpath(*parts).with_suffix(".py"))
        if top != "src":
            candidates.append((ROOT / "src").joinpath(*parts))
            candidates.append((ROOT / "src").joinpath(*parts).with_suffix(".py"))

    found: list[str] = []
    seen: set[str] = set()
    for c in candidates:
        try:
            c = c.resolve()
        except Exception:
            continue
        if not str(c).startswith(str(ROOT)):
            continue
        paths: list[Path] = []
        if c.is_file() and c.suffix == ".py":
            paths.append(c)
        elif c.is_dir():
            init = c / "__init__.py"
            if init.is_file():
                paths.append(init)
        for p in paths:
            rel = p.relative_to(ROOT).as_posix()
            if rel not in seen:
                seen.add(rel)
                found.append(rel)
    return found


def is_project_import(mod: str) -> bool:
    if not mod:
        return False
    top = mod.split(".")[0]
    if top in PROJECT_TOPS:
        return True
    return bool(resolve_module_to_files(mod))


def extract_imports(tree: ast.AST, file_path: Path) -> list[str]:
    refs: list[str] = []
    seen: set[str] = set()

    def add_mod(mod: str) -> None:
        if not mod or not is_project_import(mod):
            return
        files = resolve_module_to_files(mod)
        if files:
            for f in files:
                if f not in seen:
                    seen.add(f)
                    refs.append(f)
        elif mod not in seen:
            seen.add(mod)
            refs.append(mod)

    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                add_mod(alias.name)
        elif isinstance(node, ast.ImportFrom):
            if node.level and node.level > 0:
                rel_base = file_path.parent
                for _ in range(node.level - 1):
                    rel_base = rel_base.parent
                if node.module:
                    target = rel_base.joinpath(*node.module.split("."))
                else:
                    target = rel_base
                for c in (target.with_suffix(".py"), target / "__init__.py"):
                    if c.is_file():
                        rel = c.resolve().relative_to(ROOT).as_posix()
                        if rel not in seen:
                            seen.add(rel)
                            refs.append(rel)
                if node.module is None:
                    for alias in node.names:
                        if alias.name == "*":
                            continue
                        sub = rel_base / alias.name
                        for c in (sub.with_suffix(".py"), sub / "__init__.py"):
                            if c.is_file():
                                rel = c.resolve().relative_to(ROOT).as_posix()
                                if rel not in seen:
                                    seen.add(rel)
                                    refs.append(rel)
                elif node.module:
                    pkg_parts = list(file_path.relative_to(ROOT).parts[:-1])
                    up = node.level
                    base_parts = (
                        pkg_parts[: max(0, len(pkg_parts) - (up - 1))]
                        if up > 0
                        else pkg_parts
                    )
                    add_mod(".".join(base_parts + node.module.split(".")))
                continue
            if node.module:
                add_mod(node.module)
                for alias in node.names:
                    if alias.name == "*":
                        continue
                    sub = f"{node.module}.{alias.name}"
                    for f in resolve_module_to_files(sub):
                        if f not in seen:
                            seen.add(f)
                            refs.append(f)
    return refs


def extract_summary(source: str, tree: Optional[ast.AST], rel: str) -> str:
    doc = ast.get_docstring(tree) if tree is not None else None
    if doc:
        s = first_meaningful_summary(doc)
        if s:
            return s
    comments: list[str] = []
    for ln in source.splitlines()[:40]:
        st = ln.strip()
        if st.startswith("#!") or st.startswith("# -*-") or st.startswith("# coding"):
            continue
        if st.startswith("#"):
            comments.append(st.lstrip("# ").strip())
        elif st == "":
            if comments:
                break
        else:
            break
    if comments:
        s = first_meaningful_summary(" ".join(comments))
        if s:
            return s
    if tree is not None:
        ad = try_argparse_description(tree)
        if ad:
            return ad
    return summarize_from_filename(rel)


def main() -> None:
    rows: list[tuple[str, str, str]] = []
    errors: list[tuple[str, str]] = []
    py_files = sorted(SCRIPTS.rglob("*.py"))

    for fp in py_files:
        rel = fp.relative_to(ROOT).as_posix()
        try:
            source = fp.read_text(encoding="utf-8", errors="replace")
        except Exception as e:  # noqa: BLE001
            errors.append((rel, str(e)))
            rows.append((rel, f"(unreadable: {e})", ""))
            continue

        tree = None
        try:
            tree = ast.parse(source, filename=str(fp))
        except SyntaxError as e:
            errors.append((rel, f"syntax: {e}"))
            m = re.search(r'"""(.*?)"""', source, re.S)
            summary = (
                first_meaningful_summary(m.group(1))
                if m
                else summarize_from_filename(rel)
            )
            rows.append(
                (
                    rel,
                    summary or summarize_from_filename(rel),
                    "(parse failed — imports unavailable)",
                )
            )
            continue

        summary = extract_summary(source, tree, rel)
        refs = [r for r in extract_imports(tree, fp) if r != rel]
        rows.append((rel, summary, "; ".join(refs) if refs else ""))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scripts Analysis"

    headers = ["File Name", "Summary of functionality", "Referred files"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for col in range(1, 4):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    alt = PatternFill("solid", fgColor="F2F2F2")
    for i, (name, summary, refs) in enumerate(rows, 2):
        display = name[len("scripts/") :] if name.startswith("scripts/") else name
        ws.cell(i, 1, display)
        ws.cell(i, 2, summary)
        ws.cell(i, 3, refs)
        for col in range(1, 4):
            cell = ws.cell(i, col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin
            if i % 2 == 0:
                cell.fill = alt

    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 70
    ws.row_dimensions[1].height = 22
    ws.auto_filter.ref = f"A1:C{len(rows) + 1}"
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Counts")
    ws2.append(["Metric", "Value"])
    ws2.append(["Total Python files under scripts/", len(rows)])
    ws2.append(
        [
            "Files with referred project imports",
            sum(1 for _, _, r in rows if r and not r.startswith("(")),
        ]
    )
    ws2.append(["Files with empty referred files", sum(1 for _, _, r in rows if not r)])
    ws2.append(["Parse/read errors", len(errors)])
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)
    ws2.column_dimensions["A"].width = 45
    ws2.column_dimensions["B"].width = 15

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Rows: {len(rows)}")
    print(f"Errors: {len(errors)}")
    for r in rows[:3]:
        print("---")
        print(r[0])
        print(r[1][:220])
        print(r[2][:220] if r[2] else "(no refs)")


if __name__ == "__main__":
    main()
