"""Contract floor for the Semantic File Identity workbook enrichment.

All three workbooks are UNTRACKED (regenerable from their existing generator scripts + this
repo's disk state), so every test here ``pytest.skip``s if its workbook is absent rather than
failing — the same reasoning the Semantic OS applies to ``graph.dot`` staleness.

Covers:
  * new headers are APPENDED, not inserted — columns A-C (A-F on the tests workbook) unchanged;
  * the positional consumer contract (``grok/book_file_coverage_report.py``) still holds;
  * the join rate meets the floor on the scripts/src workbooks;
  * re-running the enricher is idempotent (never duplicates or reorders columns);
  * the tests workbook gets a coverage POINTER, never its own Semantic ID column (decision 3);
  * the "Coverage Join Method" label is honest (sentinel -> NONE; a resolved id -> IMPORT_PATH_*).
"""
from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

import pytest

_REPO = Path(__file__).resolve().parents[1]
if str(_REPO / "scripts" / "governance") not in sys.path:
    sys.path.insert(0, str(_REPO / "scripts" / "governance"))
if str(_REPO / "grok") not in sys.path:
    sys.path.insert(0, str(_REPO / "grok"))

import enrich_workbooks_with_semantic_identity as enrich  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

_WB_SCRIPTS = _REPO / "scripts_business_functionality.xlsx"
_WB_SRC = _REPO / "results" / "analysis" / "src_business_functionality.xlsx"
_WB_TESTS = _REPO / "docs" / "analysis" / "tests_functionality_inventory.xlsx"
_PROJECTION = _REPO / "data" / "semantic_os"

pytestmark = pytest.mark.skipif(openpyxl is None, reason="openpyxl not installed")


def _skip_if_absent(path: Path) -> None:
    if not path.is_file():
        pytest.skip(f"{path.relative_to(_REPO).as_posix()} is untracked/disposable and absent")


def _skip_if_no_projection() -> None:
    if not (_PROJECTION / "file_identities.jsonl").is_file():
        pytest.skip("data/semantic_os/file_identities.jsonl absent — run seed_semantic_os.py first")


# ── header / column-append discipline ──────────────────────────────────────────────────────

def test_scripts_workbook_headers_appended_not_inserted():
    _skip_if_absent(_WB_SCRIPTS)
    wb = openpyxl.load_workbook(_WB_SCRIPTS)
    ws = wb["Scripts Analysis"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers[:3] == ["File Name", "Summary of functionality", "Referred files"]
    for h in enrich._IDENTITY_HEADERS:
        assert h in headers


def test_src_workbook_headers_appended_not_inserted():
    _skip_if_absent(_WB_SRC)
    wb = openpyxl.load_workbook(_WB_SRC)
    ws = wb["src_py_inventory"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers[:3] == ["File Name", "Summary of functionality", "Referred files"]
    for h in enrich._IDENTITY_HEADERS:
        assert h in headers


def test_tests_workbook_headers_appended_not_inserted():
    _skip_if_absent(_WB_TESTS)
    wb = openpyxl.load_workbook(_WB_TESTS)
    for sheet_name in ("Test Functionality", "By File"):
        ws = wb[sheet_name]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert headers[:6] == [
            "File Name", "Relative Path", "Summary of functionality", "Referred files",
            headers[4], headers[5],  # "Test Class"/"Test Classes" differ by sheet — position-checked
        ]
        for h in enrich._COVERS_HEADERS:
            assert h in headers


def test_tests_workbook_has_no_own_identity_column():
    """Decision 3: test files get a coverage POINTER, never their own Semantic ID."""
    _skip_if_absent(_WB_TESTS)
    wb = openpyxl.load_workbook(_WB_TESTS)
    for sheet_name in ("Test Functionality", "By File"):
        ws = wb[sheet_name]
        headers = {ws.cell(1, c).value for c in range(1, ws.max_column + 1)}
        assert "Semantic ID" not in headers, f"{sheet_name} must not carry its own Semantic ID column"
        assert "Semantic Name" not in headers


# ── positional consumer contract ────────────────────────────────────────────────────────────

def test_positional_consumer_contract_holds():
    """grok/book_file_coverage_report.py reads row[0..2] only — appended columns must not break it."""
    _skip_if_absent(_WB_SRC)
    _skip_if_absent(_WB_SCRIPTS)
    import book_file_coverage_report as bfc

    src_rows = bfc.load_src_inventory()
    assert src_rows and src_rows[0]["file"].startswith("src/")
    assert all({"file", "summary", "referred"} <= set(r) for r in src_rows[:5])

    script_rows = bfc.load_scripts_inventory()
    assert script_rows and script_rows[0]["file"].startswith("scripts/")
    assert all({"file", "summary", "referred"} <= set(r) for r in script_rows[:5])


# ── join rate ────────────────────────────────────────────────────────────────────────────────

def test_join_rate_meets_floor_on_scripts_and_src():
    _skip_if_absent(_WB_SCRIPTS)
    _skip_if_absent(_WB_SRC)
    _skip_if_no_projection()
    identity_by_path = enrich.load_identity_by_path(_PROJECTION)

    for path, sheet in ((_WB_SCRIPTS, "Scripts Analysis"), (_WB_SRC, "src_py_inventory")):
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet]
        headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        sem_col = headers["Semantic ID"]
        total = 0
        joined = 0
        for r in range(2, ws.max_row + 1):
            if not ws.cell(r, 1).value:
                continue
            total += 1
            if ws.cell(r, sem_col).value:
                joined += 1
        assert total > 0
        assert joined / total >= 0.98, f"{path.name}: join rate {joined}/{total} below floor"


# ── idempotence ──────────────────────────────────────────────────────────────────────────────

def test_enrichment_is_idempotent_on_a_copy():
    """Never mutate the real artifact from a test — copy to tmp_path, run twice, compare."""
    _skip_if_absent(_WB_SCRIPTS)
    _skip_if_no_projection()
    identity_by_path = enrich.load_identity_by_path(_PROJECTION)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / _WB_SCRIPTS.name
        shutil.copy(_WB_SCRIPTS, tmp_path)

        enrich.enrich_scripts_workbook(tmp_path, identity_by_path)
        wb1 = openpyxl.load_workbook(tmp_path)
        ws1 = wb1["Scripts Analysis"]
        snapshot1 = [
            [ws1.cell(r, c).value for c in range(1, ws1.max_column + 1)]
            for r in range(1, ws1.max_row + 1)
        ]
        max_col_1 = ws1.max_column

        enrich.enrich_scripts_workbook(tmp_path, identity_by_path)
        wb2 = openpyxl.load_workbook(tmp_path)
        ws2 = wb2["Scripts Analysis"]
        snapshot2 = [
            [ws2.cell(r, c).value for c in range(1, ws2.max_column + 1)]
            for r in range(1, ws2.max_row + 1)
        ]
        max_col_2 = ws2.max_column

        assert max_col_1 == max_col_2, "column count grew on rerun — enrichment is not idempotent"
        assert snapshot1 == snapshot2, "cell values changed on a no-op rerun"


# ── join-method honesty ─────────────────────────────────────────────────────────────────────

def test_covers_semantic_id_join_method_is_honest():
    _skip_if_absent(_WB_TESTS)
    wb = openpyxl.load_workbook(_WB_TESTS)
    ws = wb["Test Functionality"]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    referred_col = headers["Referred files"]
    covers_col = headers["Covers Semantic ID"]
    method_col = headers["Coverage Join Method"]

    checked = 0
    for r in range(2, ws.max_row + 1):
        if not ws.cell(r, 1).value:
            continue
        referred = ws.cell(r, referred_col).value
        covers = ws.cell(r, covers_col).value
        method = ws.cell(r, method_col).value
        if (referred or "").strip() in enrich._REFERRED_EMPTY_SENTINELS:
            assert method == "NONE", f"row {r}: empty referred but method={method!r}"
            assert not covers
        if covers:
            assert method.startswith("IMPORT_PATH_"), f"row {r}: covers={covers!r} but method={method!r}"
        checked += 1
    assert checked > 0


def test_covers_from_referred_is_pure_and_deterministic():
    """Unit-level check of the join function itself (no workbook needed)."""
    _skip_if_no_projection()
    identity_by_path = enrich.load_identity_by_path(_PROJECTION)
    sample_path = next(iter(identity_by_path))
    matches, method = enrich.covers_from_referred(sample_path, identity_by_path)
    assert method == "IMPORT_PATH_EXACT"
    assert matches and matches[0]["physical_path"] == sample_path

    matches2, method2 = enrich.covers_from_referred("", identity_by_path)
    assert matches2 == [] and method2 == "NONE"

    matches3, method3 = enrich.covers_from_referred("src/definitely/not/real.py", identity_by_path)
    assert matches3 == [] and method3 == "UNRESOLVED:1"
