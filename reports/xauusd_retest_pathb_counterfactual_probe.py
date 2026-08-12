#!/usr/bin/env python3
"""
OBSERVE-ONLY probe — path B counterfactual on the two XAUUSD RETESTs that died
to off_session:OFF_SESSION (semantic window Jul 22 → Aug 7 2026 Excel corpus).

Does NOT modify production code, config, or ACTIVE_VERSION.
Force-session-pass is local to this process via dataclasses.replace on CRTConfig.

Targets (broker timestamps, confirmed prior turn):
  1) 2026-07-22 19:00 RETEST SHORT → soft-conf bar 19:15 rejected OFF_SESSION
  2) 2026-07-28 05:30 RETEST LONG  → soft-conf bar 05:45 rejected OFF_SESSION

For each: force session allow → capture path-A TRADE_OPENED (or not) → at the
soft-conf bar run EngineRunner → ExecutionPlanner → compute_crt_levels → UltronRiskGate.

Output: reports/xauusd_retest_pathb_counterfactual.json (+ .md summary).
"""
from __future__ import annotations

import dataclasses
import json
import os
import sys
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
os.chdir(ROOT)

from openpyxl import load_workbook

from config_layer.crt_engine_v2 import CRTEngine, Candle
from config_layer.execution_planner import ExecutionPlannerV1_2
from config_layer.production_config import (
    get_active_version,
    get_prod_metadata,
    load_prod_config_from_registry,
)
from core.decision_engine import DecisionEngine
from core.engine_runner import EngineRunner
from core.gate_intelligence import compute_crt_levels
from core.ultron_risk_gate import UltronRiskGate
from features.feature_pipeline import FeaturePipeline
from features.feature_schema import CANONICAL_FEATURES
from runtime.backtest_v2 import HTFBuilder


XLSX = ROOT / "data" / "XAUUSD_M15_20260807_203705.xlsx"
OUT_JSON = ROOT / "reports" / "xauusd_retest_pathb_counterfactual.json"
OUT_MD = ROOT / "reports" / "xauusd_retest_pathb_counterfactual.md"

# Soft-conf (reject) bars — where FILTER_REJECTED fired under production session policy
TARGETS = (
    {
        "id": "retest_1_short",
        "retest_ts": "2026-07-22 19:00:00",
        "soft_conf_ts": "2026-07-22 19:15:00",
        "expected_dir": "SHORT",
    },
    {
        "id": "retest_2_long",
        "retest_ts": "2026-07-28 05:30:00",
        "soft_conf_ts": "2026-07-28 05:45:00",
        "expected_dir": "LONG",
    },
)


def _load_candles() -> list[Candle]:
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    rows = list(wb["candles"].iter_rows(min_row=2, values_only=True))
    wb.close()
    candles: list[Candle] = []
    for i, r in enumerate(rows):
        ts, o, h, l, c, v = r
        if isinstance(ts, str):
            ts = datetime.fromisoformat(ts)
        candles.append(
            Candle(
                timestamp=ts,
                open=float(o),
                high=float(h),
                low=float(l),
                close=float(c),
                volume=float(v or 0),
                index=i,
            )
        )
    return candles


def _candles_to_df(candles: list[Candle]) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "timestamp": [c.timestamp for c in candles],
            "open": [c.open for c in candles],
            "high": [c.high for c in candles],
            "low": [c.low for c in candles],
            "close": [c.close for c in candles],
            "volume": [c.volume for c in candles],
        }
    )


def _build_feature_index(candles: list[Candle]) -> dict[str, dict[str, float]]:
    """timestamp str → feature dict (canonical + OHLCV).

    Keeps ALL input timestamps (no finalize dropna) so early soft-conf bars
    remain addressable. Warmup NaNs are ffill/bfill/0-filled — observe-only.
    """
    import numpy as np

    df = _candles_to_df(candles)
    pipe = FeaturePipeline(df)
    # Run every stage of FeaturePipeline.run except finalize's dropna.
    pipe.compute_price_features()
    pipe.compute_volume_features()
    pipe.compute_indicators()
    pipe.compute_trend_features()
    pipe.compute_volatility_regime()
    pipe.compute_context()
    pipe.compute_structure_liquidity()
    pipe.compute_normalization()
    pipe.compute_canonical_price_features()
    pipe.compute_canonical_volatility_features()
    pipe.compute_canonical_ema_features()
    pipe.compute_canonical_trend_features()
    pipe.compute_canonical_structure_features()
    pipe.compute_canonical_temporal_features()
    pipe.compute_liquidity_distance()
    pipe.promote_volume_spike()
    pipe.compute_canonical_session()

    out_df = pipe.df.replace([np.inf, -np.inf], np.nan).copy()
    for col in CANONICAL_FEATURES:
        if col in out_df.columns:
            out_df[col] = out_df[col].ffill().bfill().fillna(0.0)
        else:
            out_df[col] = 0.0
    pipe.df = out_df
    vectors = pipe.build_feature_vector()
    arr = np.asarray(vectors)
    if arr.ndim == 1:
        arr = arr.reshape(1, -1)

    ts_col = "timestamp" if "timestamp" in out_df.columns else out_df.columns[0]
    index: dict[str, dict[str, float]] = {}
    n = min(len(out_df), len(arr))
    for i in range(n):
        ts = out_df.iloc[i][ts_col]
        key = pd.Timestamp(ts).strftime("%Y-%m-%d %H:%M:%S")
        feat = {
            CANONICAL_FEATURES[j]: float(arr[i, j])
            for j in range(min(len(CANONICAL_FEATURES), arr.shape[1]))
        }
        for k in ("open", "high", "low", "close", "volume"):
            if k in out_df.columns:
                feat[k] = float(out_df.iloc[i][k])
            elif k not in feat:
                feat[k] = 0.0
        # FeaturePipeline "atr" slot is relative; planner/levels need price units.
        atr_slot = float(feat.get("atr") or 0.0)
        close = float(feat.get("close") or 0.0)
        if 0 < atr_slot < 1.0 and close > 0:
            # relative ATR × close → absolute
            feat["atr"] = atr_slot * close
            feat["atr_relative"] = atr_slot
        elif atr_slot <= 0:
            feat["atr"] = max(float(feat["high"]) - float(feat["low"]), 1e-6)
        index[key] = feat
    return index


def _merge_engine_config(prod: dict, *, force_session_pass: bool = True) -> dict:
    """Mirror live_engine_hook._load_engine_config flatten (read-only process copy)."""
    engine_cfg = dict(prod["engine_runner"])
    decision_cfg = dict(prod["decision_engine"])
    if force_session_pass:
        # Path-B force-session-pass: allow every canonical session name the
        # adapter recognises so OFF_SESSION/ASIA/CLOSED bars are not adapter-killed.
        # Process-local only — production JSON is never written.
        engine_cfg["allowed_sessions"] = [
            "asia", "london", "new_york", "overlap", "closed", "off_session",
        ]
    merged: dict[str, Any] = {}
    merged.update(engine_cfg)
    merged.update(decision_cfg)
    merged["fusion_engine"] = dict(prod["fusion_engine"])
    merged["ultron_risk_gate"] = dict(prod["ultron_risk_gate"])
    merged["execution_planner"] = dict(prod["execution_planner"])
    merged["crt_engine"] = dict(prod["crt_engine"])
    if "rr_fusion" not in merged and "rr_fusion" in engine_cfg:
        merged["rr_fusion"] = engine_cfg["rr_fusion"]
    return merged


def _dir_int(direction_name: str | None) -> int:
    if direction_name == "LONG":
        return 1
    if direction_name == "SHORT":
        return -1
    return 0


def _safe_json(obj: Any) -> Any:
    if obj is None or isinstance(obj, (str, int, float, bool)):
        return obj
    if isinstance(obj, dict):
        return {str(k): _safe_json(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_safe_json(v) for v in obj]
    if hasattr(obj, "value"):
        return obj.value
    if hasattr(obj, "name"):
        return obj.name
    try:
        return float(obj)
    except Exception:
        return str(obj)


def run_path_a_force_session(candles: list[Candle]) -> dict[str, Any]:
    """CRT replay with allowed_sessions expanded so OFF_SESSION passes."""
    base_cfg = load_prod_config_from_registry(get_active_version(), "XAUUSD")
    # Force-session-pass: include every named window + OFF_SESSION.
    forced = dataclasses.replace(
        base_cfg,
        allowed_sessions=("LONDON", "NEWYORK", "OVERLAP", "ASIA", "OFF_SESSION"),
    )
    engine = CRTEngine(forced)
    htf = HTFBuilder(4, "XAUUSD")
    initialized = False

    target_soft = {t["soft_conf_ts"] for t in TARGETS}
    target_retest = {t["retest_ts"] for t in TARGETS}
    events: list[dict] = []
    trade_opened: dict[str, dict] = {}
    soft_conf_outcomes: dict[str, dict] = {}
    inverted_sl: list[dict] = []

    orig_record = engine.ev_log.record

    def record(event, candle, **kwargs):
        if event in ("FILTER_REJECTED", "TRADE_OPENED", "RETEST_CONFIRMED", "CONFIRMATION_FAILED") or "RETEST" in str(event):
            events.append(
                {
                    "event": event,
                    "ts": str(candle.timestamp),
                    "reason": kwargs.get("reason"),
                    "metadata": _safe_json(kwargs.get("metadata")),
                    "direction": _safe_json(kwargs.get("direction")),
                    "price": kwargs.get("price"),
                }
            )
        return orig_record(event, candle, **kwargs)

    engine.ev_log.record = record  # type: ignore[method-assign]

    # Wrap build_trade to capture inverted-SL geometry when it returns None
    orig_build = engine.executor.build_trade

    def build_trade_traced(state, risk_engine=None):
        trade = orig_build(state, risk_engine)
        # Snapshot geometry even on failure
        disp = state.displacement_candle
        retest = state.retest_candle
        atr = float(state.atr_abs or 0.0)
        buf = float(forced.sl_atr_buffer)
        dir_name = state.direction.name if state.direction else None
        entry = float(retest.close) if retest else None
        sl_would = None
        if disp is not None and atr > 0 and dir_name in ("LONG", "SHORT"):
            if dir_name == "LONG":
                sl_would = float(disp.low) - buf * atr
            else:
                sl_would = float(disp.high) + buf * atr
        inverted = False
        if entry is not None and sl_would is not None:
            if dir_name == "LONG" and sl_would >= entry:
                inverted = True
            if dir_name == "SHORT" and sl_would <= entry:
                inverted = True
        snap = {
            "ts": str(retest.timestamp) if retest else None,
            "dir": dir_name,
            "entry": entry,
            "sl_would": sl_would,
            "disp_high": float(disp.high) if disp else None,
            "disp_low": float(disp.low) if disp else None,
            "disp_open": float(disp.open) if disp else None,
            "disp_close": float(disp.close) if disp else None,
            "atr": atr,
            "sl_atr_buffer": buf,
            "inverted_sl": inverted,
            "trade_built": trade is not None,
        }
        if inverted or trade is None:
            inverted_sl.append(snap)
        return trade

    engine.executor.build_trade = build_trade_traced  # type: ignore[method-assign]

    for candle in candles:
        completed = htf.push(candle)
        if not initialized:
            if completed:
                engine.initialise_range(htf.seed_candles(), htf.current_htf_id, "UNKNOWN")
                initialized = True
            else:
                continue

        before = engine.state.current_state.name
        action = engine.process_candle(candle, htf.current_htf_id)
        after = engine.state.current_state.name
        act = action.get("action") if isinstance(action, dict) else str(action)
        ts = candle.timestamp.strftime("%Y-%m-%d %H:%M:%S")

        if ts in target_retest or ts in target_soft or act in (
            "TRADE_OPENED", "FILTER_REJECTED", "RETEST_CONFIRMED", "CONFIRMATION_FAILED"
        ):
            st = engine.state
            trade = st.active_trade
            row = {
                "ts": ts,
                "state_before": before,
                "state_after": after,
                "action": act,
                "direction": st.direction.name if st.direction else None,
                "has_active_trade": trade is not None and getattr(trade, "status", None) in ("OPEN", "TP1"),
                "soft_conf_candles": getattr(st, "soft_conf_candles", None),
            }
            if trade is not None and act == "TRADE_OPENED":
                row["trade"] = {
                    "entry": float(trade.entry_price),
                    "sl": float(trade.sl_price),
                    "tp1": float(trade.tp1_price),
                    "tp2": float(trade.tp2_price),
                    "risk_pct": float(trade.risk_pct),
                    "direction": trade.direction.name if hasattr(trade.direction, "name") else str(trade.direction),
                    "id": trade.id,
                }
                trade_opened[ts] = row["trade"]
            if ts in target_soft:
                soft_conf_outcomes[ts] = row
            events.append({"event": "ACTION", **row})

    return {
        "mode": "path_a_force_session",
        "allowed_sessions_forced": list(forced.allowed_sessions),
        "trade_opened_by_ts": trade_opened,
        "soft_conf_outcomes": soft_conf_outcomes,
        "build_trade_snaps": inverted_sl,
        "events": events,
    }


def run_path_b_at_bars(
    candles: list[Candle],
    feat_index: dict[str, dict[str, float]],
    path_a: dict[str, Any],
) -> list[dict[str, Any]]:
    prod = get_prod_metadata()
    eng_cfg = _merge_engine_config(prod)
    runner = EngineRunner(eng_cfg)
    planner = ExecutionPlannerV1_2(prod["execution_planner"])
    ultron = UltronRiskGate(prod["ultron_risk_gate"])
    crt_section = prod["crt_engine"]

    results: list[dict[str, Any]] = []
    for tgt in TARGETS:
        soft_ts = tgt["soft_conf_ts"]
        retest_ts = tgt["retest_ts"]
        out: dict[str, Any] = {
            "id": tgt["id"],
            "retest_ts": retest_ts,
            "soft_conf_ts": soft_ts,
            "expected_dir": tgt["expected_dir"],
            "path_a_trade_at_soft_or_retest": None,
            "engine_runner": None,
            "planner": None,
            "crt_levels": None,
            "ultron": None,
            "verdict": None,
            "errors": [],
        }

        # Path A trade if force-session opened on soft-conf or retest bar
        pa_trades = path_a.get("trade_opened_by_ts") or {}
        trade_a = pa_trades.get(soft_ts) or pa_trades.get(retest_ts)
        out["path_a_trade_at_soft_or_retest"] = trade_a

        # Direction: prefer path-A trade, else expected
        if trade_a:
            dir_name = trade_a.get("direction") or tgt["expected_dir"]
        else:
            dir_name = tgt["expected_dir"]
        d_int = _dir_int(dir_name)

        feat = feat_index.get(soft_ts)
        if feat is None:
            out["errors"].append(f"no feature vector for {soft_ts}")
            out["verdict"] = "INCOMPLETE_NO_FEATURES"
            results.append(out)
            continue

        # EngineRunner input: features + OHLCV + integrity + CRT direction
        input_data = dict(feat)
        input_data["_data_integrity"] = "real"
        input_data["instrument"] = "XAUUSD"
        input_data["direction"] = d_int
        input_data["signal_dir"] = d_int
        input_data["trade_direction"] = d_int
        # atr may already be in features under different name
        if "atr" not in input_data or not input_data.get("atr"):
            # atr_relative * close if present
            atr_rel = float(input_data.get("atr_relative") or 0.0)
            close = float(input_data.get("close") or 0.0)
            input_data["atr"] = atr_rel * close if atr_rel and close else float(input_data.get("atr") or 1.0)

        context = {
            "instrument": "XAUUSD",
            "timeframe": "M15",
            "strategy_consensus_direction": d_int,
            "symbol": "XAUUSD",
        }

        try:
            er = runner.run(input_data, context=context)
            out["engine_runner"] = {
                "decision": er.get("decision") or er.get("status"),
                "reason": er.get("reason"),
                "reject_stage": er.get("reject_stage"),
                "final_score": er.get("final_score") or er.get("score"),
                "selected_direction": er.get("selected_direction"),
                "confidence": er.get("confidence"),
                "regime": er.get("regime"),
                "engine_scores": _safe_json(er.get("engine_scores") or er.get("scores")),
                "keys": sorted(str(k) for k in er.keys())[:40],
            }
        except Exception as exc:
            out["errors"].append(f"EngineRunner: {exc}")
            out["engine_runner"] = {"error": str(exc), "trace": traceback.format_exc()[-800:]}
            out["verdict"] = "ENGINE_RUNNER_ERROR"
            results.append(out)
            continue

        er_decision = str((out["engine_runner"] or {}).get("decision") or "").lower()
        execute_tokens = {"execute", "accept", "go", "approved", "approve"}
        sel_dir = er.get("selected_direction")
        if sel_dir is None:
            sel_dir = d_int
        try:
            sel_dir = int(sel_dir)
        except Exception:
            sel_dir = d_int

        # ── Secondary observe arm ────────────────────────────────────────────
        # EngineRunner builds zone_gate_ctx as zone_result.get("passed"), but
        # zone_result only carries passed inside meta (engine_runner.py:712-717,
        # 1019-1021). When DecisionEngine rejects zone_gate_invalid, re-evaluate
        # DecisionEngine with valid=meta.passed so we can measure the NEXT gates
        # (low_score / planner / Ultron) without mutating production code.
        zone_bypass_used = False
        if er_decision not in execute_tokens and str(er.get("reason", "")) == "zone_gate_invalid":
            try:
                # Re-run engines is expensive; instead re-score DecisionEngine with
                # the same final_score/p_win and zone valid forced True (meta.passed
                # was true in collector logs for these bars).
                de = DecisionEngine(eng_cfg)
                final_score = float(er.get("final_score") or 0.0)
                p_win = float(er.get("confidence") or final_score or 0.5)
                # weak_component = 1 - final_score (mirrors EngineRunner fusion_ctx)
                re_eval = de.evaluate(
                    score=final_score,
                    p_win=p_win if p_win > 0 else 0.5,
                    zone_gate={"valid": True, "score": final_score},
                    fusion={
                        "normalized_score": final_score,
                        "weak_component": max(0.0, 1.0 - final_score),
                        "zone_gate_dead": False,
                    },
                    config=eng_cfg,
                )
                out["decision_reeval_zone_valid"] = _safe_json(re_eval)
                if str(re_eval.get("decision", "")).lower() == "execute":
                    zone_bypass_used = True
                    er = {**er, **re_eval, "decision": "execute", "reason": "zone_valid_reeval_observe_only"}
                    er_decision = "execute"
                    out["engine_runner"]["zone_valid_reeval"] = "execute"
                    out["engine_runner"]["zone_valid_reeval_reason"] = re_eval.get("reason")
                else:
                    out["engine_runner"]["zone_valid_reeval"] = re_eval.get("decision")
                    out["engine_runner"]["zone_valid_reeval_reason"] = re_eval.get("reason")
            except Exception as exc:
                out["errors"].append(f"zone_valid_reeval: {exc}")

        if er_decision not in execute_tokens:
            reval = out.get("decision_reeval_zone_valid") or {}
            reval_reason = reval.get("reason") if isinstance(reval, dict) else None
            primary = er.get("reason") or er_decision or "empty"
            if reval_reason:
                out["verdict"] = (
                    f"PATH_B_STOP_ENGINE_RUNNER:{primary}"
                    f"|zone_valid_reeval:{reval_reason}"
                )
            else:
                out["verdict"] = f"PATH_B_STOP_ENGINE_RUNNER:{primary}"
            results.append(out)
            continue

        # Planner
        try:
            engine_result = {
                "decision": "execute",
                "selected_direction": sel_dir if sel_dir != 0 else d_int,
                "confidence": float(er.get("confidence") or er.get("final_score") or 0.0),
                "regime": er.get("regime") or "unknown",
            }
            out["zone_bypass_used"] = zone_bypass_used
            plan_features = dict(feat)
            for k in ("open", "high", "low", "close", "volume", "atr"):
                plan_features.setdefault(k, input_data.get(k, 0.0))
            # common planner feature aliases
            for alias_src, alias_dst in (
                ("body_ratio", "body_ratio"),
                ("disp_strength", "disp_strength"),
                ("retest_depth", "retest_depth"),
                ("atr_relative", "atr_relative"),
            ):
                if alias_src in feat:
                    plan_features[alias_dst] = feat[alias_src]
            ctx = {
                "symbol": "XAUUSD",
                "signal": engine_result["selected_direction"],
                "score": float(engine_result["confidence"]),
                "account_balance": 100_000.0,
            }
            plan = planner.plan(engine_result, plan_features, ctx)
            out["planner"] = _safe_json(plan)
        except Exception as exc:
            out["errors"].append(f"Planner: {exc}")
            out["planner"] = {"error": str(exc), "trace": traceback.format_exc()[-800:]}
            out["verdict"] = "PLANNER_ERROR"
            results.append(out)
            continue

        if str(plan.get("decision", "")).lower() != "execute":
            out["verdict"] = f"PATH_B_STOP_PLANNER:{plan.get('decision')}"
            results.append(out)
            continue

        # CRT levels (live injection)
        try:
            intent = str(plan.get("trade_intent", "UNKNOWN")).lower()
            tp1_key = f"tp1_atr_multiplier_{intent}"
            tp1_mult = float(crt_section.get(tp1_key, crt_section.get("tp1_atr_multiplier", 1.0)))
            tp2_mult = float(crt_section.get("tp2_atr_multiplier", 2.0))
            entry = float(plan["entry_price"])
            direction = int(plan.get("direction", d_int))
            levels = compute_crt_levels(
                entry=entry,
                direction=direction,
                low=float(plan_features["low"]),
                high=float(plan_features["high"]),
                atr=float(plan_features.get("atr") or input_data["atr"]),
                sl_atr_buffer=float(crt_section["sl_atr_buffer"]),
                tp1_mult=tp1_mult,
                tp2_mult=tp2_mult,
            )
            plan["stop_loss"] = levels["sl"]
            plan["take_profit_1"] = levels["tp1"]
            plan["take_profit_2"] = levels["tp2"]
            risk_px = abs(entry - levels["sl"])
            plan["rr_ratio"] = round(abs(levels["tp1"] - entry) / risk_px, 6) if risk_px > 0 else 0.0
            plan["rr_source"] = "sl_tp_geometry"
            plan["risk_percent"] = float(prod["execution_planner"]["risk_percent"])
            out["crt_levels"] = _safe_json(levels)
            out["planner_with_levels"] = {
                "entry": entry,
                "sl": plan["stop_loss"],
                "tp1": plan["take_profit_1"],
                "tp2": plan["take_profit_2"],
                "rr_ratio": plan["rr_ratio"],
                "intent": plan.get("trade_intent"),
                "direction": direction,
            }
        except Exception as exc:
            out["errors"].append(f"compute_crt_levels: {exc}")
            out["verdict"] = "CRT_LEVELS_ERROR"
            results.append(out)
            continue

        # Ultron
        try:
            portfolio_state = {
                "daily_trade_count": 0,
                "daily_loss_pct": 0.0,
                "open_risk_pct": 0.0,
                "account_balance": 100_000.0,
                "timestamp": soft_ts,
            }
            # trade dict expected by Ultron
            trade_dict = dict(plan)
            # common key aliases
            trade_dict.setdefault("entry", plan["entry_price"])
            trade_dict.setdefault("entry_price", plan["entry_price"])
            trade_dict.setdefault("stop_loss", plan["stop_loss"])
            trade_dict.setdefault("take_profit", plan["take_profit_1"])
            ug = ultron.evaluate(trade_dict, portfolio_state)
            out["ultron"] = _safe_json(ug)
            u_dec = str(ug.get("decision", "")).upper()
            if u_dec == "APPROVE":
                out["verdict"] = "PATH_B_WOULD_OPEN"
            else:
                out["verdict"] = f"PATH_B_STOP_ULTRON:{u_dec}:{ug.get('risk_reason') or ug.get('reason')}"
        except Exception as exc:
            out["errors"].append(f"Ultron: {exc}")
            out["ultron"] = {"error": str(exc), "trace": traceback.format_exc()[-800:]}
            out["verdict"] = "ULTRON_ERROR"

        results.append(out)
    return results


def _write_md(report: dict) -> None:
    lines = [
        "# XAUUSD RETEST path-B counterfactual (observe-only)",
        "",
        f"**Generated:** {report['generated_at']}",
        f"**ACTIVE_VERSION:** `{report['active_version']}`",
        f"**Corpus:** `{report['corpus']}`",
        f"**HTF:** {report['htf_candles_per_range']}",
        "",
        "> Force-session-pass is process-local only (`allowed_sessions` includes `OFF_SESSION`).",
        "> No production config or engine code was modified.",
        "",
        "## Path A (force session)",
        "",
        f"- Forced allow list: `{report['path_a']['allowed_sessions_forced']}`",
        f"- TRADE_OPENED timestamps: `{list((report['path_a'].get('trade_opened_by_ts') or {}).keys())}`",
        f"- Soft-conf outcomes: `{json.dumps(report['path_a'].get('soft_conf_outcomes') or {}, default=str)}`",
        "",
        "### build_trade geometry snaps (incl. inverted SL)",
        "",
        "```json",
        json.dumps(report["path_a"].get("build_trade_snaps") or [], indent=2, default=str),
        "```",
        "",
    ]
    for t_ts, trade in (report["path_a"].get("trade_opened_by_ts") or {}).items():
        lines.append(f"### TRADE_OPENED @ `{t_ts}`")
        lines.append("")
        lines.append("```json")
        lines.append(json.dumps(trade, indent=2))
        lines.append("```")
        lines.append("")

    lines += ["## Path B (EngineRunner → Planner → CRT levels → Ultron)", ""]
    for row in report["path_b"]:
        lines.append(f"### {row['id']} — soft-conf `{row['soft_conf_ts']}`")
        lines.append("")
        lines.append(f"- **Verdict:** `{row.get('verdict')}`")
        lines.append(f"- Expected dir: `{row.get('expected_dir')}`")
        lines.append(f"- Path-A trade present: `{bool(row.get('path_a_trade_at_soft_or_retest'))}`")
        er = row.get("engine_runner") or {}
        lines.append(
            f"- EngineRunner: decision=`{er.get('decision')}` reason=`{er.get('reason')}` "
            f"stage=`{er.get('reject_stage')}` score=`{er.get('final_score')}`"
        )
        pl = row.get("planner") or {}
        lines.append(f"- Planner: decision=`{pl.get('decision') if isinstance(pl, dict) else pl}`")
        if row.get("planner_with_levels"):
            lines.append(f"- Levels: `{json.dumps(row['planner_with_levels'])}`")
        ug = row.get("ultron") or {}
        if isinstance(ug, dict):
            lines.append(
                f"- Ultron: decision=`{ug.get('decision')}` reason=`{ug.get('risk_reason') or ug.get('reason')}`"
            )
        if row.get("errors"):
            lines.append(f"- Errors: `{row['errors']}`")
        lines.append("")

    lines += [
        "## Interpretation",
        "",
        "- **Path A force-session:** soft-conf + zone can reach `EXECUTION`; `TRADE_OPENED` still "
        "requires `build_trade` non-inverted SL.",
        "- **Path B force-session (production EngineRunner):** first live decision reject reason "
        "is authoritative for that call.",
        "- **`zone_valid_reeval` arm:** if `zone_gate_invalid` was only the `passed`-flag wiring "
        "(`zone_result.meta.passed` true but top-level `passed` absent), re-evaluate DecisionEngine "
        "with `valid=True` to name the *next* semantic reject (`low_score` / `weak_setup` / …).",
        "- `PATH_B_WOULD_OPEN` only if EngineRunner→Planner→levels→Ultron all approve under the "
        "empty portfolio fixture.",
        "- Observe-only: **no** session-policy or zone authority granted.",
        "",
    ]
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    print("Loading candles…")
    candles = _load_candles()
    print(f"  {len(candles)} bars  {candles[0].timestamp} → {candles[-1].timestamp}")

    print("Building feature index (FeaturePipeline)…")
    try:
        feat_index = _build_feature_index(candles)
        print(f"  feature rows: {len(feat_index)}")
    except Exception as exc:
        print(f"  FeaturePipeline failed: {exc}")
        traceback.print_exc()
        # Minimal OHLCV-only fallback (EngineRunner may reject incomplete features)
        feat_index = {}
        for c in candles:
            key = c.timestamp.strftime("%Y-%m-%d %H:%M:%S")
            feat = {name: 0.0 for name in CANONICAL_FEATURES}
            feat.update(
                {
                    "open": c.open,
                    "high": c.high,
                    "low": c.low,
                    "close": c.close,
                    "volume": c.volume,
                    "atr": max(c.high - c.low, 1e-6),
                }
            )
            feat_index[key] = feat
        print(f"  fallback OHLCV+zeros: {len(feat_index)} rows")

    print("Path A — force session pass…")
    path_a = run_path_a_force_session(candles)
    print(f"  TRADE_OPENED: {list(path_a['trade_opened_by_ts'].keys())}")
    print(f"  soft_conf_outcomes: {path_a.get('soft_conf_outcomes')}")
    print(f"  build_trade_snaps: {path_a.get('build_trade_snaps')}")

    print("Path B — EngineRunner / Planner / Ultron…")
    path_b = run_path_b_at_bars(candles, feat_index, path_a)
    for row in path_b:
        print(f"  {row['id']}: {row.get('verdict')}")
        er = row.get("engine_runner") or {}
        print(f"    ER decision={er.get('decision')} reason={er.get('reason')} score={er.get('final_score')}")

    report = {
        "generated_at": datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
        "active_version": get_active_version(),
        "corpus": str(XLSX.relative_to(ROOT)),
        "htf_candles_per_range": 4,
        "observe_only": True,
        "production_behavior_changed": False,
        "targets": list(TARGETS),
        "path_a": {
            "allowed_sessions_forced": path_a["allowed_sessions_forced"],
            "trade_opened_by_ts": path_a["trade_opened_by_ts"],
            "soft_conf_outcomes": path_a.get("soft_conf_outcomes"),
            "build_trade_snaps": path_a.get("build_trade_snaps"),
            "events": path_a["events"],
        },
        "path_b": path_b,
    }
    OUT_JSON.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    _write_md(report)
    print(f"Wrote {OUT_JSON}")
    print(f"Wrote {OUT_MD}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
